import json, sys, copy, datetime
import os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from classify import classify
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as CL

OUT = sys.argv[1]
W = os.path.join(os.path.dirname(HERE), "")  # repo root: expects input/ and work/ folders
D = json.load(open(W + "work/tb_all.json"))
SRC = {"FY24": W + "input/TrialBal_FY_23-24.xlsx", "FY25": W + "input/TrialBal_FY_24-25.xlsx", "FY26": W + "input/TrialBal_FY_25-26.xlsx"}
TBN = {"FY26": "TB 2025-26", "FY25": "TB 2024-25", "FY24": "TB 2023-24"}

# ------------------------------------------------------------------ styles (reference: Times New Roman 10)
TN = "Times New Roman"
fN = Font(name=TN, size=10)
fB = Font(name=TN, size=10, bold=True)
fI = Font(name=TN, size=10, italic=True)
fBI = Font(name=TN, size=10, bold=True, italic=True)
fT = Font(name=TN, size=11, bold=True)
fS = Font(name=TN, size=8, color="808080")           # small reviewer check cells
fG = Font(name=TN, size=9, italic=True, color="808080")
fBlue = Font(name=TN, size=10, color="0000FF")
ACC = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '
ACC0 = '_ * #,##0_ ;_ * \\-#,##0_ ;_ * "-"??_ ;_ @_ '
PCT = '0.00%'
thin = Side(style="thin")
dbl = Side(style="double")
bTB = Border(top=thin, bottom=thin)
bT = Border(top=thin)
bTot = Border(top=thin, bottom=dbl)
wrapc = Alignment(horizontal="center", vertical="center", wrap_text=True)
wrapl = Alignment(horizontal="left", vertical="top", wrap_text=True)
just = Alignment(horizontal="justify", vertical="top", wrap_text=True)
yel = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()
wb.remove(wb.active)


def sheet(name, widths, grid=False):
    ws = wb.create_sheet(name)
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.sheet_view.showGridLines = grid
    return ws


def w(ws, ref, v, font=fN, fmt=None, al=None, bd=None, fill=None):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt: c.number_format = fmt
    if al: c.alignment = al
    if bd: c.border = bd
    if fill: c.fill = fill
    return c


def q(s): return "'" + s + "'"


# ------------------------------------------------------------------ MasterData
MD = "MasterData"
ws = sheet(MD, {"A": 44, "B": 30, "C": 4, "D": 6, "E": 12, "F": 28, "G": 20, "H": 12, "I": 10})
md = {}
rows = [
    ("Company Information", None, None),
    ("CIN", "U72900TN2016PTC113553", "cin"),
    ("Company Name", "TECHFRUITS SOLUTIONS LIMITED", "name"),
    ("Formerly known as", "(Formerly known as Techfruits Solutions Private Limited)", "former"),
    ("Registered Address", "Flat 1D, Kaashyap Enclave, New No.209, Old No.13A, Velachery Main Road, Velachery, Chennai - 600042", "addr"),
    ("Authorised Capital (Rs)", 1000000, "auth"),
    ("Face value per equity share (Rs)", 10, "fv"),
    ("Bonus / split factor after March 31, 2026 (1 = none)", 1, "bonus"),
    (None, None, None),
    ("RFS Details", None, None),
    ("Year ended (latest)", datetime.date(2026, 3, 31), "d26"),
    ("Year ended (previous)", datetime.date(2025, 3, 31), "d25"),
    ("Year ended (earliest)", datetime.date(2024, 3, 31), "d24"),
    ("Opening balance sheet date", datetime.date(2023, 3, 31), "d23"),
    ("Scale", 100000, "scale"),
    ("Closing stock-in-trade as at March 31, 2026 (Rs)", 11015077, "clstk"),
    ("Tax rate - Sec 115BAA (22% + 10% SC + 4% cess)", "=22%*1.1*1.04", "tax"),
    ("Gratuity (prior period) relating to FY 2023-24 (Rs)", 0, "grat24"),
    (None, None, None),
    ("Audit Firm Name", "[Peer reviewed auditor]", "firm"),
    ("Firm Reg No", "[FRN]", "frn"),
    ("Peer Review Certificate No", "[Certificate No.]", "peer"),
    ("Engagement Partner Name", "[Partner]", "partner"),
    ("Membership No", "[M.No.]", "mno"),
    ("Date of Signing the Examination Report", "xx/xx/2026", "sdate"),
    ("UDIN : ", "XXXXXXXXXX", "udin"),
    ("Place", "Chennai", "place"),
]
r = 1
for lbl, val, key in rows:
    if lbl:
        w(ws, f"A{r}", lbl, fB if key is None else fN)
    if key:
        c = w(ws, f"B{r}", val, fBlue)
        if isinstance(val, datetime.date): c.number_format = "dd-mm-yyyy"
        elif key == "tax": c.number_format = "0.000%"
        elif isinstance(val, (int, float)): c.number_format = "#,##0"
        if key in ("clstk", "grat24", "bonus", "firm", "frn", "peer", "partner", "mno", "sdate", "udin"): c.fill = yel
        md[key] = f"{MD}!$B${r}"
    r += 1
w(ws, "B15", 100000, fBlue, "#,##0")
w(ws, "C15", None)
w(ws, "F15", "All amounts in Indian Rupees Lakhs, except as otherwise stated", fI)
w(ws, "C16", None)
w(ws, "F16", "As per audited financial statements FY 2025-26 (not carried in Tally TB)", fI)
w(ws, "F18", "Balance treated as relating to periods up to March 31, 2023 (replace with actuarial split)", fI)
# Director / signatory table (reference layout)
w(ws, "D2", "Director/Signatory Details", fB)
for i, h in enumerate(["Sr. No", "DIN", "Name", "Designation", "Category", "Signatory"]):
    w(ws, f"{CL(4+i)}3", h, fB, bd=bTB)
dirs = [(1, "07190573", "C K LENIN CHRISTOPHER", "Director", "Promoter", "Yes"),
        (2, "06934008", "W VINOLIN ANTONY JOANS", "Director", "Promoter", "Yes"),
        (3, "[DIN]", "P SURESH KUMAR", "Director", "-", "No")]
for k, d in enumerate(dirs):
    for i, v in enumerate(d):
        w(ws, f"{CL(4+i)}{4+k}", v, fBlue if i in (1, 2, 3, 4) else fN)
md["dir1"], md["dir1des"], md["dir1din"] = f"{MD}!$F$4", f"{MD}!$G$4", f"{MD}!$E$4"
md["dir2"], md["dir2des"], md["dir2din"] = f"{MD}!$F$5", f"{MD}!$G$5", f"{MD}!$E$5"
md["dir3"] = f"{MD}!$F$6"
SCALE = md["scale"]


def dtxt(key, pre="As at \n"):
    return f'="{pre}"&TEXT({md[key]},"mmmm dd, yyyy")'


# ------------------------------------------------------------------ labels per code
LBL = {
    "SC": ("Equity share capital", "Cr", "BS"), "RS": ("Reserves and surplus", "Cr", "BS"),
    "LTB_SEC": ("Term loans from banks and NBFCs (secured)", "Cr", "BS"),
    "DTL": ("Deferred tax liability (net)", "Cr", "BS"), "OLTL": ("Rental and security deposits received", "Cr", "BS"),
    "LTP_GRAT": ("Provision for gratuity", "Cr", "BS"),
    "STB_BANK": ("Overdraft facilities from banks (secured)", "Cr", "BS"),
    "STB_OTH": ("Bill discounting and channel finance (secured)", "Cr", "BS"),
    "TP_MSME": ("Dues to micro and small enterprises", "Cr", "BS"), "TP_OTH": ("Dues to other creditors", "Cr", "BS"),
    "OCL_STAT": ("Statutory dues payable", "Cr", "BS"), "OCL_EMP": ("Salary and other employee dues payable", "Cr", "BS"),
    "OCL_EXP": ("Expenses and other payables", "Cr", "BS"), "OCL_ADV": ("Advance from customers", "Cr", "BS"),
    "STP_TAX": ("Provision for income tax (net of advance tax)", "Cr", "BS"),
    "PPE_LAND": ("Land", "Dr", "BS"), "PPE_BLDG": ("Buildings", "Dr", "BS"), "PPE_FURN": ("Furniture and fittings", "Dr", "BS"),
    "PPE_VEH": ("Vehicles", "Dr", "BS"), "PPE_OFFEQ": ("Office equipment", "Dr", "BS"), "PPE_ELEC": ("Electrical installations", "Dr", "BS"),
    "PPE_COMP": ("Computers", "Dr", "BS"), "INT_SW": ("Computer software", "Dr", "BS"), "CWIP": ("Capital work-in-progress", "Dr", "BS"),
    "IAUD": ("Intangible assets under development", "Dr", "BS"), "DTA": ("Deferred tax asset (net)", "Dr", "BS"),
    "LTLA": ("Security deposits", "Dr", "BS"), "TR": ("Trade receivables", "Dr", "BS"),
    "CASH": ("Cash on hand", "Dr", "BS"), "BANK": ("Balances with banks in current accounts", "Dr", "BS"),
    "BANK_FD": ("Fixed deposits with banks", "Dr", "BS"),
    "STLA_SUP": ("Advance to suppliers", "Dr", "BS"), "STLA_EMP": ("Advances to employees", "Dr", "BS"), "STLA_OTH": ("Other advances", "Dr", "BS"),
    "OCA_GOVT": ("Balances with government authorities", "Dr", "BS"), "OCA_TAX": ("Advance tax, TDS and TCS receivable", "Dr", "BS"),
    "OCA_PRE": ("Prepaid expenses", "Dr", "BS"), "OCA_UNB": ("Unbilled revenue", "Dr", "BS"), "OCA_EMD": ("Earnest money deposits", "Dr", "BS"),
    "OCA_INT": ("Interest accrued on deposits", "Dr", "BS"), "OCA_IPO": ("IPO expenses pending adjustment", "Dr", "BS"),
    "OCA_OTH": ("Other current assets", "Dr", "BS"),
    "XX_OPSTK": ("Opening stock", "Dr", "BS"), "XX_PLA": ("Net profit as per books", "Dr", "BS"), "XX_BR": ("Inter-branch accounts", "Dr", "BS"),
    "REV_PROD": ("Sale of products", "Cr", "P&L"), "REV_SERV": ("Sale of services", "Cr", "P&L"), "REV_OTH": ("Other operating revenue", "Cr", "P&L"),
    "OI_INT": ("Interest income", "Cr", "P&L"), "OI_OTH": ("Other non-operating income", "Cr", "P&L"),
    "PUR": ("Purchases of stock-in-trade", "Dr", "P&L"), "DIR": ("Direct expenses", "Dr", "P&L"),
    "EMP_SAL": ("Salaries, wages and incentives", "Dr", "P&L"), "EMP_DIR": ("Directors' remuneration", "Dr", "P&L"),
    "EMP_PF": ("Contribution to provident and other funds", "Dr", "P&L"), "EMP_GRAT": ("Gratuity", "Dr", "P&L"),
    "EMP_WEL": ("Staff welfare expenses", "Dr", "P&L"),
    "FIN_INT": ("Interest on borrowings", "Dr", "P&L"), "FIN_OTH": ("Other borrowing costs", "Dr", "P&L"),
    "DEP": ("Depreciation and amortisation", "Dr", "P&L"),
    "OE_POWER": ("Electricity charges", "Dr", "P&L"), "OE_RENT": ("Rent", "Dr", "P&L"), "OE_RATES": ("Rates and taxes", "Dr", "P&L"),
    "OE_INS": ("Insurance", "Dr", "P&L"), "OE_RM": ("Repairs and maintenance", "Dr", "P&L"), "OE_TRAV": ("Travelling and conveyance", "Dr", "P&L"),
    "OE_COMM": ("Brokerage and commission", "Dr", "P&L"), "OE_ADV": ("Advertisement and business promotion", "Dr", "P&L"),
    "OE_PROF": ("Legal and professional charges", "Dr", "P&L"), "OE_AUD": ("Payment to auditors", "Dr", "P&L"),
    "OE_TEL": ("Communication, postage and courier", "Dr", "P&L"), "OE_PRINT": ("Printing and stationery", "Dr", "P&L"),
    "OE_FRT": ("Freight and transportation", "Dr", "P&L"), "OE_BAD": ("Bad debts and balances written off", "Dr", "P&L"),
    "OE_FOREX": ("Loss on foreign exchange fluctuation", "Dr", "P&L"), "OE_MISC": ("Miscellaneous expenses", "Dr", "P&L"),
    "TAX_CUR": ("Current tax", "Dr", "P&L"), "TAX_DEF": ("Deferred tax", "Dr", "P&L"),
}
SPECIAL = {
    "Lenin (Director)": ("Loan from director - C K Lenin Christopher", "Cr", "BS"),
    "Vinolin (Director)": ("Loan from director - W Vinolin Antony Joans", "Cr", "BS"),
    "P Suresh Kumar": ("Loan from director - P Suresh Kumar", "Cr", "BS"),
    "Interest on MSME Payable": ("Interest accrued on dues to MSME", "Cr", "BS"),
    "Interest on MSME": ("Interest on delayed payments to MSME", "Dr", "P&L"),
    "Directors Remuneration Payable": ("Directors' remuneration payable", "Cr", "BS"),
}
CRED_PRIM = {"Capital Account", "Loans (Liability)", "Current Liabilities", "Sales Accounts", "Indirect Incomes"}


def master_for(x, which):
    """which = 3 (closing) or 0 (opening)"""
    v = x["v"][which]
    other = x["v"][0 if which == 3 else 3]
    s = v if abs(v) > 0.001 else (other if abs(other) > 0.001 else (-1 if x["primary"] in CRED_PRIM else 1))
    dr, cr = classify(x["primary"], x["parent"], x["name"])
    code = dr if s >= 0 else cr
    if x["name"] in SPECIAL and x["primary"] in ("Loans (Liability)", "Current Liabilities", "Indirect Expenses"):
        if not (x["name"] == "Directors Remuneration Payable" and s > 0):
            return SPECIAL[x["name"]]
    if code in ("OCL_EXP",) and ("credit card" in x["name"].lower()):
        return ("Credit card dues", "Cr", "BS")
    return LBL[code]


# ------------------------------------------------------------------ TB sheets (replica of Tally export + Masters)
used_labels = set()
TBROW = {}
for fy in ["FY26", "FY25", "FY24"]:
    src = openpyxl.load_workbook(SRC[fy]).active
    ws = wb.create_sheet(TBN[fy])
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style: continue
            n = ws.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                n.font = copy.copy(c.font); n.number_format = c.number_format
                n.alignment = copy.copy(c.alignment); n.border = copy.copy(c.border); n.fill = copy.copy(c.fill)
    for m in src.merged_cells.ranges: ws.merge_cells(str(m))
    for k, d in src.column_dimensions.items(): ws.column_dimensions[k].width = d.width
    ws.column_dimensions["F"].width = 7; ws.column_dimensions["G"].width = 44
    ws.column_dimensions["H"].width = 15; ws.column_dimensions["I"].width = 12
    hf = Font(name="Arial", size=9, bold=True)
    for col, h in zip("FGHI", ["Group", "Masters", "Net Amount", "In Lakhs"]):
        c = ws[f"{col}7"]; c.value = h; c.font = hf; c.border = Border(bottom=thin)
    ws["I6"].value = f"={SCALE}"; ws["I6"].font = Font(name="Arial", size=9); ws["I6"].number_format = "#,##0"
    if fy == "FY24":
        for col, h in zip("JKL", ["Masters (Opening)", "Opening Net Amount", "Opening In Lakhs"]):
            c = ws[f"{col}7"]; c.value = h; c.font = hf; c.border = Border(bottom=thin)
        ws.column_dimensions["J"].width = 44; ws.column_dimensions["K"].width = 15; ws.column_dimensions["L"].width = 12
    df = Font(name="Arial", size=9)
    for x in D[fy]:
        r = x["r"]
        if x["type"] == "X":
            c = ws[f"G{r}"]; c.value = "Repeated in Tally export - not considered"; c.font = Font(name="Arial", size=8, italic=True, color="808080")
            continue
        if not x["inc"]: continue
        lbl, nat, grp = master_for(x, 3)
        used_labels.add(lbl)
        sgn = 1 if x["v"][3] >= 0 else -1
        flip = (sgn > 0 and nat == "Cr") or (sgn < 0 and nat == "Dr")
        ws[f"F{r}"].value = grp; ws[f"G{r}"].value = lbl
        ws[f"H{r}"].value = f"=-E{r}" if flip and abs(x["v"][3]) > 0.001 else f"=E{r}"
        ws[f"I{r}"].value = f"=H{r}/$I$6"
        for col in "FGHI": ws[f"{col}{r}"].font = df
        ws[f"H{r}"].number_format = "#,##0.00"; ws[f"I{r}"].number_format = "#,##0.00"
        if fy == "FY24":
            l2, n2, g2 = master_for(x, 0)
            used_labels.add(l2)
            s2 = 1 if x["v"][0] >= 0 else -1
            f2 = (s2 > 0 and n2 == "Cr") or (s2 < 0 and n2 == "Dr")
            ws[f"J{r}"].value = l2
            ws[f"K{r}"].value = f"=-B{r}" if f2 and abs(x["v"][0]) > 0.001 else f"=B{r}"
            ws[f"L{r}"].value = f"=K{r}/$I$6"
            for col in "JKL": ws[f"{col}{r}"].font = df
            ws[f"K{r}"].number_format = "#,##0.00"; ws[f"L{r}"].number_format = "#,##0.00"
        if x["name"] not in TBROW.get(fy, {}):
            TBROW.setdefault(fy, {})[x["name"]] = r


def tbsum(fy, label_ref, col="I", mcol="G"):
    s = q(TBN[fy])
    return f"SUMIF({s}!${mcol}:${mcol},{label_ref},{s}!${col}:${col})"


# ------------------------------------------------------------------ Adj Entry RFS (journal, Rs. lakhs)
ADJ = "Adj Entry RFS"
wsA = sheet(ADJ, {"A": 12, "B": 70, "C": 40, "D": 8, "E": 14, "F": 14})
w(wsA, "A1", "Restatement adjustment entries (Rs. in lakhs)", fB)
for col, h in zip("ABCDEF", ["Date", "Particulars", "Masters", "Group", "Dr", "Cr"]):
    w(wsA, f"{col}2", h, fB, bd=bTB, al=wrapc)
R_S = TBROW["FY25"]["Reserves & Surplus"]
w(wsA, "H2", "Working", fB)
w(wsA, "H3", "Prior period gratuity debited to reserves in FY 2024-25 (TB 2024-25, Reserves & Surplus - Debit)")
w(wsA, "L3", f"={q(TBN['FY25'])}!C{R_S}/{SCALE}", fmt=ACC)
w(wsA, "H4", "Out of above relating to FY 2023-24 (MasterData)")
w(wsA, "L4", f"={md['grat24']}/{SCALE}", fmt=ACC)
w(wsA, "H5", "Relating to periods up to March 31, 2023")
w(wsA, "L5", "=L3-L4", fmt=ACC)
w(wsA, "H6", "Tax rate")
w(wsA, "L6", f"={md['tax']}", fmt="0.000%")
entries = [
    (datetime.date(2023, 3, 31), [("Reserves and surplus", "Reserves and surplus", "BS", "=L5", None),
                                  ("   To Provision for gratuity", "Provision for gratuity", "BS", None, "=E{r0}")],
     "(Being gratuity liability relating to periods up to March 31, 2023, booked in FY 2024-25 as prior period item, now recognised in the opening balance sheet)"),
    (datetime.date(2023, 3, 31), [("Deferred tax asset (net)", "Deferred tax asset (net)", "BS", "=ROUND(L5*L6,2)", None),
                                  ("   To Reserves and surplus", "Reserves and surplus", "BS", None, "=E{r0}")],
     "(Being deferred tax asset on the above opening gratuity liability)"),
    (datetime.date(2024, 3, 31), [("Gratuity (Employee benefits expense - FY 2023-24)", "Gratuity", "P&L", "=L4", None),
                                  ("   To Provision for gratuity", "Provision for gratuity", "BS", None, "=E{r0}")],
     "(Being gratuity expense relating to FY 2023-24 charged to that year)"),
    (datetime.date(2024, 3, 31), [("Deferred tax asset (net)", "Deferred tax asset (net)", "BS", "=ROUND(L4*L6,2)", None),
                                  ("   To Deferred tax (P&L - FY 2023-24)", "Deferred tax", "P&L", None, "=E{r0}")],
     "(Being deferred tax credit on gratuity expense of FY 2023-24)"),
    (datetime.date(2025, 3, 31), [("Provision for gratuity", "Provision for gratuity", "BS", "=L3", None),
                                  ("   To Reserves and surplus", "Reserves and surplus", "BS", None, "=E{r0}")],
     "(Being reversal of prior period gratuity debited directly to reserves in the books for FY 2024-25, already recognised in earlier years above)"),
    (datetime.date(2025, 3, 31), [("Deferred tax (P&L - FY 2024-25)", "Deferred tax", "P&L", "=ROUND(L3*L6,2)", None),
                                  ("   To Deferred tax asset (net)", "Deferred tax asset (net)", "BS", None, "=E{r0}")],
     "(Being reversal of deferred tax asset recognised in earlier years, as the audited accounts for FY 2024-25 carry the full liability)"),
]
r = 3
for dt, legs, narr in entries:
    r0 = r
    for i, (p, m, g, dr, cr) in enumerate(legs):
        if i == 0:
            c = w(wsA, f"A{r}", dt, fmt="dd-mm-yyyy")
        else:
            c = w(wsA, f"A{r}", f"=A{r0}", fmt="dd-mm-yyyy")
        w(wsA, f"B{r}", p)
        w(wsA, f"C{r}", m, fG); w(wsA, f"D{r}", g, fG)
        if dr: w(wsA, f"E{r}", dr.replace("{r0}", str(r0)), fmt=ACC)
        if cr: w(wsA, f"F{r}", cr.replace("{r0}", str(r0)), fmt=ACC)
        r += 1
    w(wsA, f"B{r}", narr, fI, al=wrapl)
    wsA.row_dimensions[r].height = 26
    r += 2
ADJ_LAST = r + 20
w(wsA, f"B{r}", "Further adjustments, if any, may be added above this line using the same Masters names as in BS_1 / PL_2.", fG)
AR = lambda col: f"{q(ADJ)}!${col}$3:${col}${ADJ_LAST}"


def adj_bs(label_ref, date_key, nat):
    dr = f'SUMIFS({AR("E")},{AR("C")},{label_ref},{AR("A")},"<="&{md[date_key]})'
    cr = f'SUMIFS({AR("F")},{AR("C")},{label_ref},{AR("A")},"<="&{md[date_key]})'
    return f"={cr}-{dr}" if nat == "Cr" else f"={dr}-{cr}"


def adj_rs(label_ref, date_key):
    own = f'SUMIFS({AR("F")},{AR("C")},{label_ref},{AR("A")},"<="&{md[date_key]})-SUMIFS({AR("E")},{AR("C")},{label_ref},{AR("A")},"<="&{md[date_key]})'
    pl = f'SUMIFS({AR("E")},{AR("D")},"P&L",{AR("A")},"<="&{md[date_key]})-SUMIFS({AR("F")},{AR("D")},"P&L",{AR("A")},"<="&{md[date_key]})'
    return f"={own}-({pl})"


def adj_pl(label_ref, end_key, start_key, nat):
    crit = f'{AR("C")},{label_ref},{AR("A")},">"&{md[start_key]},{AR("A")},"<="&{md[end_key]}'
    dr = f'SUMIFS({AR("E")},{crit})'; cr = f'SUMIFS({AR("F")},{crit})'
    return f"={dr}-{cr}" if nat == "Dr" else f"={cr}-{dr}"


# ------------------------------------------------------------------ common header / signature
def head(ws, title, units=True, col="B", span="F"):
    w(ws, f"{col}2", f"={md['name']}", fB); w(ws, f"{col}3", f"={md['former']}", fB); w(ws, f"{col}4", f'="CIN: "&{md["cin"]}', fB)
    w(ws, f"{col}5", title, fB)
    if units: w(ws, f"{col}6", "(All amounts in Rs. lakhs, except as otherwise stated)", fI)
    for rr in (2, 3, 4, 5):
        ws.merge_cells(f"{col}{rr}:{span}{rr}")
        ws[f"{col}{rr}"].alignment = Alignment(horizontal="center")
    if units:
        ws.merge_cells(f"{col}6:{span}6"); ws[f"{col}6"].alignment = Alignment(horizontal="center")


def sign_block(ws, r, cB="B", cD="D", cF="F"):
    w(ws, f"{cB}{r}", "As per our report of even date attached"); r += 2
    w(ws, f"{cB}{r}", f'="for "&{md["firm"]}', fB); w(ws, f"{cD}{r}", "for and on behalf of the Board of Directors of", fI); r += 1
    w(ws, f"{cB}{r}", "Chartered Accountants", fI); w(ws, f"{cD}{r}", f"={md['name']}", fB); r += 1
    w(ws, f"{cB}{r}", f'="Firm registration number: "&{md["frn"]}'); w(ws, f"{cD}{r}", f'="CIN: "&{md["cin"]}'); r += 1
    w(ws, f"{cB}{r}", f'="Peer review certificate number: "&{md["peer"]}'); r += 3
    for c in (cB, cD, cF): w(ws, f"{c}{r}", "Sd/-")
    r += 1
    w(ws, f"{cB}{r}", f"={md['partner']}", fB); w(ws, f"{cD}{r}", f"=PROPER({md['dir1']})", fB); w(ws, f"{cF}{r}", f"=PROPER({md['dir2']})", fB); r += 1
    w(ws, f"{cB}{r}", "Partner", fI); w(ws, f"{cD}{r}", f"={md['dir1des']}", fI); w(ws, f"{cF}{r}", f"={md['dir2des']}", fI); r += 1
    w(ws, f"{cB}{r}", f'="Membership number: "&{md["mno"]}'); w(ws, f"{cD}{r}", f'="DIN : "&{md["dir1din"]}'); w(ws, f"{cF}{r}", f'="DIN : "&{md["dir2din"]}'); r += 1
    w(ws, f"{cB}{r}", f'="UDIN : "&{md["udin"]}'); r += 2
    for c in (cB, cD, cF): w(ws, f"{c}{r}", f'="Place: "&{md["place"]}')
    r += 1
    for c in (cB, cD, cF): w(ws, f"{c}{r}", f'="Date: "&{md["sdate"]}')
    return r


# ================================================================== SAP_4 (background & policies)
SAP = "SAP_4"
wsS = sheet(SAP, {"A": 1, "B": 5, "C": 110})
head(wsS, "ANNEXURE IV - BACKGROUND, BASIS OF PREPARATION AND SIGNIFICANT ACCOUNTING POLICIES", units=False, span="C")
w(wsS, "B8", 1, fB); w(wsS, "C8", "BACKGROUND", fB)
bg = ("M/s. Techfruits Solutions Limited (\"the Company\") was incorporated on [date] as a private limited company under the Companies Act, 2013 with the Registrar of "
      "Companies, Chennai (CIN U72900TN2016PTC113553). The Company was converted into a public limited company and the name was changed to Techfruits Solutions Limited "
      "pursuant to a fresh certificate of incorporation dated [date]. The registered office of the Company is at Velachery, Chennai, Tamil Nadu.")
w(wsS, "C10", bg, al=just); wsS.row_dimensions[10].height = 54
w(wsS, "C11", "Main objects of the Company:\nThe Company is engaged in the business of trading in IT hardware and software, system integration, networking and related "
              "services including annual maintenance contracts. The Company operates from Tamil Nadu and has branches in Kerala, Maharashtra and Andhra Pradesh.", al=just)
wsS.row_dimensions[11].height = 48
w(wsS, "B12", "=B8+1", fB); w(wsS, "C12", "SIGNIFICANT ACCOUNTING POLICIES", fB)
pol = [
    "a) Basis of preparation\nThe Restated Summary Statement of Assets and Liabilities of the Company as at March 31, 2026, March 31, 2025 and March 31, 2024 and the related Restated "
    "Summary Statements of Profit and Loss and Cash Flows for the years then ended, together with the annexures thereto (collectively, the \"Restated Summary Statements\"), have been "
    "prepared specifically for inclusion in the Draft Red Herring Prospectus / Red Herring Prospectus / Prospectus to be filed by the Company with the SME Platform of the Stock "
    "Exchange in connection with its proposed Initial Public Offer. The Restated Summary Statements have been prepared to comply with Section 26 of Part I of Chapter III of the "
    "Companies Act, 2013, the Securities and Exchange Board of India (Issue of Capital and Disclosure Requirements) Regulations, 2018, as amended, and the Guidance Note on Reports "
    "in Company Prospectuses (Revised 2019) issued by the Institute of Chartered Accountants of India.",
    "The Restated Summary Statements have been compiled from the audited financial statements of the Company for the years ended March 31, 2026, March 31, 2025 and March 31, 2024, "
    "prepared under the historical cost convention on accrual basis in accordance with the Accounting Standards specified under Section 133 of the Companies Act, 2013 read with the "
    "Companies (Accounting Standards) Rules, 2021, and are presented in the format prescribed in Division I of Schedule III to the Act.",
    "Operating Cycle:\nAll assets and liabilities have been classified as current or non-current as per the Company's normal operating cycle and other criteria set out in Schedule III. "
    "Based on the nature of products and the time between acquisition of assets and their realisation in cash, the Company has ascertained its operating cycle as twelve months.",
    "b) Use of estimates\nThe preparation of the financial statements in conformity with Indian GAAP requires the management to make estimates and assumptions that affect the reported "
    "amounts of assets and liabilities, disclosure of contingent liabilities and the reported income and expenses during the year. Differences between actual results and estimates are "
    "recognised in the period in which the results are known / materialise.",
    "c) Revenue recognition\nRevenue from sale of goods is recognised when the significant risks and rewards of ownership are transferred to the buyer, which generally coincides with "
    "delivery. Revenue from services, including annual maintenance contracts, is recognised as the services are rendered, on a proportionate basis over the period of the contract. "
    "Services rendered but not billed at the year end are shown as unbilled revenue. Revenue is stated net of GST, trade discounts and returns. Interest income is recognised on a "
    "time proportion basis.",
    "d) Cash and bank balances\nCash and cash equivalents comprise cash in hand and balances in current accounts with banks. Fixed deposits with banks, including those held as margin "
    "money, are disclosed as other bank balances.",
    "e) Cash flow statement\nCash flows are reported using the indirect method, whereby profit before tax is adjusted for the effects of transactions of a non-cash nature and any "
    "deferrals or accruals of past or future cash receipts or payments.",
    "f) Property, Plant and Equipment and Intangible assets\nProperty, Plant and Equipment are stated at cost less accumulated depreciation. Cost comprises the purchase price and any "
    "attributable cost of bringing the asset to its working condition for its intended use. Intangible assets are stated at cost less accumulated amortisation. Expenditure on "
    "development of software meeting the recognition criteria of AS 26 is carried as intangible assets under development until the asset is ready for use.",
    "g) Depreciation\nDepreciation on Property, Plant and Equipment is provided on the written down value method based on the useful lives prescribed in Schedule II to the Companies "
    "Act, 2013. Land is not depreciated. Computer software is amortised over its estimated useful life.",
    "h) Inventories\nStock-in-trade is valued at lower of cost and net realisable value. Cost includes purchase price and other costs incurred in bringing the inventories to their present "
    "location and condition.",
    "i) Employee benefits\nShort term employee benefits are recognised as an expense in the year in which the related service is rendered. Contributions to provident fund and ESI are "
    "charged to the Statement of Profit and Loss. Liability for gratuity, a defined benefit obligation, is provided on the basis of valuation as at the balance sheet date.",
    "j) Foreign currency transactions\nTransactions in foreign currency are recorded at the exchange rate prevailing on the date of the transaction. Monetary items outstanding at the "
    "year end are restated at the closing rate and the resultant differences are recognised in the Statement of Profit and Loss.",
    "k) Taxes on income\nCurrent tax is determined as the amount of tax payable in respect of taxable income for the year under the Income-tax Act, 1961; the Company has opted for "
    "the concessional regime under Section 115BAA. Deferred tax is recognised on timing differences, subject to consideration of prudence, using tax rates enacted or substantively "
    "enacted at the balance sheet date.",
    "l) Leases\nLease rentals in respect of assets taken on operating lease are charged to the Statement of Profit and Loss over the lease term.",
    "m) Borrowing costs\nBorrowing costs attributable to the acquisition of qualifying assets are capitalised as part of the cost of such assets. All other borrowing costs are charged "
    "to the Statement of Profit and Loss.",
    "n) Provisions and contingencies\nA provision is recognised when the Company has a present obligation as a result of a past event, it is probable that an outflow of resources will "
    "be required to settle the obligation and a reliable estimate can be made. Contingent liabilities are disclosed in the notes. Contingent assets are not recognised.",
    "o) Earnings per share\nBasic earnings per share is computed by dividing the restated profit after tax by the weighted average number of equity shares outstanding during the year, "
    "adjusted for bonus issue / share split, if any, including those occurring after the balance sheet date. There are no dilutive potential equity shares.",
    "p) Segment reporting\nThe Company is engaged in a single business segment of IT products and related services and operates primarily in India. Accordingly, no separate segment "
    "information is disclosed.",
]
r = 13
for t in pol:
    w(wsS, f"C{r}", t, al=just)
    wsS.row_dimensions[r].height = max(30, 14 * (len(t) // 135 + 1 + t.count("\n")))
    r += 1

# ================================================================== Notes (linked to TB with SUMIF on Masters, as in reference)
NE, NL, NP, NA, NPL = "Notes_Equity & Reserves", "Notes_Liabilities", "Notes_PPE", "Notes_Assets", "Notes_P&L"
BS, PL, CF, SA = "BS_1", "PL_2", "CFS_3", "SOA_5"
RPT, GC, TX, AR_, AN, CP, AD = "RPT", "Gratuity and contingencies", "Tax Computation", "Accounting ratios", "Analytical Ratios", "Capitalisation", "Additional Disclosures"
SR, DRA, CRA = "Shareholder Register", "Debtor Ageing", "Creditor Ageing"
INCOME = {"Sale of products", "Sale of services", "Other operating revenue", "Interest income", "Other non-operating income"}
FYL = ["FY26", "FY25", "FY24"]
PREVFY = {"FY26": "FY25", "FY25": "FY24"}
END = {"FY26": "d26", "FY25": "d25", "FY24": "d24"}
START = {"FY26": "d25", "FY25": "d24", "FY24": "d23"}


def nz(f): return f[1:] if f.startswith("=") else f


def tbf(fy, ref, col="I", mcol="G"):
    s = q(TBN[fy])
    return f"SUMIF({s}!${mcol}:${mcol},{ref},{s}!${col}:${col})"


def tb_open(ref):   # balance as at March 31, 2023 from opening columns of TB 2023-24
    return tbf("FY24", ref, "L", "J")


NR = {}     # master label -> (sheet, row)
NTOT = {}   # head -> (sheet, row)


def note_sheet(name, span="G"):
    ws = sheet(name, {"A": 1, "B": 5, "C": 60, "D": 16.8, "E": 16.8, "F": 16.8, "G": 15, "H": 3, "I": 12, "J": 12, "K": 12})
    head(ws, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", span="F")
    return ws


class Note:
    def __init__(self, ws, first_no, r=8):
        self.ws, self.r, self.nextno, self.no = ws, r, first_no, {}

    def title(self, key, text, pl=False):
        ws, r = self.ws, self.r
        w(ws, f"B{r}", self.nextno, fB); w(ws, f"C{r}", text, fB)
        self.no[key] = f"{q(ws.title)}!$B${r}"
        self.nextno = f"=B{r}+1"
        self.pl = pl
        r += 2
        w(ws, f"C{r}", "Particulars", fB, bd=bTB)
        pre = "For the year ended\n" if pl else "As at \n"
        for i, k in enumerate(["d26", "d25", "d24"]):
            w(ws, f"{'DEF'[i]}{r}", dtxt(k, pre), fB, bd=bTB, al=wrapc)
        if not pl:
            w(ws, f"G{r}", dtxt("d23"), fG, bd=bTB, al=wrapc)
        ws.row_dimensions[r].height = 30
        self.r = r + 1

    def cols(self): return "DEF" if self.pl else "DEFG"

    def tb(self, label, nat, adj=False, show=None):
        """note line linked to TB: SUMIF(TB Masters, C, TB In Lakhs)"""
        ws, r = self.ws, self.r
        w(ws, f"C{r}", label)
        for i, fy in enumerate(FYL):
            f = "=" + tbf(fy, f"$C{r}")
            if adj:
                f += "+" + (nz(adj_pl(f"$C{r}", END[fy], START[fy], nat)) if self.pl else "(" + nz(adj_bs(f"$C{r}", END[fy], nat)) + ")")
            w(ws, f"{'DEF'[i]}{r}", f, fmt=ACC)
        if not self.pl:
            f = "=" + tb_open(f"$C{r}")
            if adj: f += "+(" + nz(adj_bs(f"$C{r}", "d23", nat)) + ")"
            w(ws, f"G{r}", f, fG, ACC)
        NR[label] = (ws.title, r)
        self.r += 1
        return r

    def line(self, label, f=None, bold=False, italic=False, fmt=ACC, bd=None, cols=None):
        ws, r = self.ws, self.r
        w(ws, f"C{r}", label, fB if bold else (fI if italic else fN), bd=bd)
        if f is not None:
            for i, col in enumerate(cols or self.cols()):
                v = f(i, col) if callable(f) else f
                if v is not None:
                    w(ws, f"{col}{r}", v, fG if col == "G" else (fB if bold else fN), fmt, bd=bd)
        self.r += 1
        return r

    def total(self, first, last, key=None, label="Total"):
        r = self.line(label, lambda i, c: f"=SUM({c}{first}:{c}{last})", bold=True, bd=bTot)
        if key: NTOT[key] = (self.ws.title, r)
        return r

    def text(self, t, h=1):
        ws, r = self.ws, self.r
        w(ws, f"C{r}", t, al=just); ws.merge_cells(f"C{r}:F{r}")
        ws.row_dimensions[r].height = 14 * h
        self.r += 1

    def gap(self, n=1): self.r += n


def nref(label, col):
    s, r = NR[label]
    return f"{q(s)}!{col}{r}"


def tref(key, col):
    s, r = NTOT[key]
    return f"{q(s)}!{col}{r}"


NOTE_REF = {}

# ---------------- Notes_Equity & Reserves
wsE = note_sheet(NE)
ne = Note(wsE, f"={SAP}!B12+1")
ne.title("SC", "Share capital")
ne.line("Authorised share capital", bold=True)
ne.line('=TEXT(' + md['auth'] + '/' + md['fv'] + ',"#,##0")&" Equity shares of Rs. "&' + md['fv'] + '&"/- each"', lambda i, c: f"={md['auth']}/{SCALE}")
ne.gap()
ne.line("Issued, subscribed and fully paid up", bold=True)
a = ne.tb("Equity share capital", "Cr")
desc_row = ne.line("", None, italic=True)
sc_tot = ne.total(a, a, "SC")
ne.gap()
ne.text("Terms / rights attached to equity shares: The Company has only one class of equity shares having a par value of Rs. 10 per share. Each holder of equity shares is "
        "entitled to one vote per share. In the event of liquidation, the equity shareholders are eligible to receive the remaining assets of the Company after distribution of all "
        "preferential amounts, in proportion to their shareholding.", 3)
ne.gap()
ne.line("(a) Reconciliation of the number of shares outstanding at the beginning and at the end of the year", bold=True)
sh_b = ne.line("No. of shares at the beginning of the year", lambda i, c: f"={'EFG'[i]}{sc_tot}*{SCALE}/{md['fv']}", fmt=ACC0, cols="DEF")
sh_i = ne.line("Add: Shares issued during the year", lambda i, c: f"={c}{sc_tot}*{SCALE}/{md['fv']}-{c}{sh_b}", fmt=ACC0, cols="DEF")
sh_c = ne.line("No. of shares at the end of the year", lambda i, c: f"=SUM({c}{sh_b}:{c}{sh_i})", bold=True, fmt=ACC0, bd=bTot, cols="DEF")
wsE[f"C{desc_row}"].value = '="("&TEXT(D' + str(sh_c) + ',"#,##0")&" Equity shares of Rs. "&' + md['fv'] + '&"/- each fully paid up)"'
NR["__shares"] = (NE, sh_c)
ne.gap()
ne.line("(b) Details of shareholding of promoters", bold=True)
r = ne.r
w(wsE, f"C{r}", "Name of promoter", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsE, f"{'DEF'[i]}{r}", dtxt(k, "No. of shares as at\n"), fB, bd=bTB, al=wrapc)
wsE.row_dimensions[r].height = 30
ne.r += 1
for k, nm in enumerate(["C K Lenin Christopher", "W Vinolin Antony Joans"]):
    rr = ne.line(f"{k+1}. {nm}", lambda i, c, k=k: f"={q(SR)}!{'CDE'[i]}{5+k}", fmt=ACC0, cols="DEF")
    ne.line("   % of total shares", lambda i, c, rr=rr: f"=IFERROR({c}{rr}/{c}{sh_c},0)", fmt=PCT, italic=True, cols="DEF")
    ne.line("   % change during the year", lambda i, c, rr=rr: (f"={c}{rr+1}-{'EF'[i]}{rr+1}" if i < 2 else "=0"), fmt=PCT, italic=True, cols="DEF")
ne.gap()
ne.line("(c) Details of shareholders holding more than 5% of shares", bold=True)
for k, nm in enumerate(["C K Lenin Christopher", "W Vinolin Antony Joans", "P Suresh Kumar"]):
    rr = ne.line(f"{k+1}. {nm} (No. of shares)", lambda i, c, k=k: f"={q(SR)}!{'CDE'[i]}{5+k}", fmt=ACC0, cols="DEF")
    ne.line("   % holding", lambda i, c, rr=rr: f"=IFERROR({c}{rr}/{c}{sh_c},0)", fmt=PCT, italic=True, cols="DEF")
ne.text("As per records of the Company, including its register of shareholders / members. There are no shares reserved for issue under options, no bonus shares, no shares "
        "allotted for consideration other than cash and no shares bought back during the five years immediately preceding March 31, 2026.", 2)
ne.gap()
ne.title("RS", "Reserves and surplus")
ne.line("Surplus in the Statement of Profit and Loss", bold=True)
RSR = {}
o = ne.line("Opening balance", lambda i, c: (f"={'EFG'[i]}{ne.r + 4}" if i < 3 else None))
p_ = ne.line("Add: Profit for the year", lambda i, c: (f"={PL}!{c}{{PAT}}" if i < 3 else None))
RSR["profit"] = p_


def book_other(i, c):
    fy = FYL[i]
    cur = tbf(fy, '"Reserves and surplus"')
    prev = tbf(PREVFY[fy], '"Reserves and surplus"') if fy in PREVFY else tb_open('"Reserves and surplus"')
    return f"={cur}-{prev}-{tbf(fy, chr(34) + 'Net profit as per books' + chr(34))}"


q1 = ne.line("Add / (Less): Adjustments made directly in reserves in the books", lambda i, c: book_other(i, c) if i < 3 else None)


def direct_adj(i, c):
    fy = FYL[i]
    crit = f'{AR("C")},"Reserves and surplus",{AR("A")},">"&{md[START[fy]]},{AR("A")},"<="&{md[END[fy]]}'
    return f"=SUMIFS({AR('F')},{crit})-SUMIFS({AR('E')},{crit})"


q2 = ne.line("Add / (Less): Restatement adjustments (Refer Annexure V)", lambda i, c: direct_adj(i, c) if i < 3 else None)
cl = ne.line("Balance at the end of the year", lambda i, c: (f"=SUM({c}{o}:{c}{q2})" if i < 3 else
                                                             "=" + tb_open('"Reserves and surplus"') + "+(" + nz(adj_rs('"Reserves and surplus"', "d23")) + ")"),
             bold=True, bd=bTot)
NTOT["RS"] = (NE, cl)
RSR.update({"open": o, "book": q1, "direct": q2, "close": cl})
bk = ne.line("Balance as per books of account", lambda i, c: "=" + (tbf(FYL[i], '"Reserves and surplus"') if i < 3 else tb_open('"Reserves and surplus"')), italic=True)
RSR["bookbal"] = bk
for i, c in enumerate("DEF"):
    w(wsE, f"{'IJK'[i]}{cl}", f"=ROUND({c}{cl}-{c}{bk}-({nz(adj_rs(chr(34)+'Reserves and surplus'+chr(34), END[FYL[i]]))}),2)", fS, ACC)
ne.text("Note: Adjustments made directly in reserves in the books for FY 2024-25 represent prior period gratuity, which has been realigned to the respective years in the "
        "Restated Summary Statements. FY 2025-26 movement represents entries passed directly to Reserves and surplus in the books.", 2)
NOTE_REF.update({"NOTE_SC": ne.no["SC"], "NOTE_RS": ne.no["RS"]})

# ---------------- Notes_Liabilities
wsL = note_sheet(NL)
nl = Note(wsL, f"={ne.no['RS']}+1")
nl.title("LTB", "Long-term borrowings")
nl.line("Secured", bold=True)
a = nl.tb("Term loans from banks and NBFCs (secured)", "Cr")
nl.line("Unsecured - Loans from directors (interest free)", bold=True)
DIRS = ["Loan from director - C K Lenin Christopher", "Loan from director - W Vinolin Antony Joans", "Loan from director - P Suresh Kumar"]
for d_ in DIRS: b_ = nl.tb(d_, "Cr")
nl.total(a, b_, "LTB")
nl.text("Loans from directors are interest free and repayable on demand at the option of the lender. As there is no stipulation for repayment within twelve months, these have been "
        "classified as long-term borrowings. Vehicle and term loans are secured by hypothecation of the respective assets.", 2)
nl.gap()
nl.line("Movement in loans from directors during the year", bold=True)
DIRMOV = {}
for kind, col_, lbl in [("acc", "D", "Loans accepted"), ("rep", "C", "Loans repaid")]:
    r = nl.r
    w(wsL, f"C{r}", lbl, fBI, bd=bTB)
    for i, k in enumerate(["d26", "d25", "d24"]):
        w(wsL, f"{'DEF'[i]}{r}", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
    wsL.row_dimensions[r].height = 30
    nl.r += 1
    for d_ in DIRS:
        r = nl.r
        w(wsL, f"C{r}", d_)
        for i, fy in enumerate(FYL):
            w(wsL, f"{'DEF'[i]}{r}", "=" + tbf(fy, f"$C{r}", col_) + f"/{SCALE}", fmt=ACC)
        DIRMOV[(kind, d_)] = r
        nl.r += 1
nl.gap()
nl.title("OLTL", "Other long-term liabilities")
a = nl.tb("Rental and security deposits received", "Cr")
nl.total(a, a, "OLTL")
nl.gap()
nl.title("LTP", "Long-term provisions")
nl.line("Provision for employee benefits", bold=True)
a = nl.tb("Provision for gratuity", "Cr", adj=True)
nl.total(a, a, "LTP")
nl.gap()
nl.title("STB", "Short-term borrowings")
nl.line("Secured", bold=True)
a = nl.tb("Overdraft facilities from banks (secured)", "Cr")
b_ = nl.tb("Bill discounting and channel finance (secured)", "Cr")
nl.total(a, b_, "STB")
nl.text("Overdraft facilities are secured by hypothecation of current assets, lien on fixed deposits and personal guarantee of the directors. [Sanction terms to be updated]", 1)
nl.gap()
nl.title("TP", "Trade payables")
a = nl.tb("Dues to micro and small enterprises", "Cr")
b_ = nl.tb("Dues to other creditors", "Cr")
nl.total(a, b_, "TP")
NTOT["TP_MSME"], NTOT["TP_OTH"] = (NL, a), (NL, b_)
nl.gap()
nl.line("Disclosure under Micro, Small and Medium Enterprises Development Act, 2006", bold=True)
MSME_ROWS = nl.r
nl.line("(a) Principal amount remaining unpaid to suppliers at the end of the year", lambda i, c: f"={c}{a}", cols="DEF")
msme_b = nl.line("(b) Interest due thereon remaining unpaid at the end of the year", None)
msme_c = nl.line("(c) Interest paid / payable for delayed payments during the year", None)
nl.line("(d) Amount of further interest remaining due and payable in succeeding years", 0, cols="DEF")
nl.text("The above information has been determined to the extent such parties have been identified on the basis of information available with the Company.", 1)
nl.gap()
AGE_TP = []
for k, (dk, base) in enumerate([("d26", 5), ("d25", 13), ("d24", 21)]):
    r = nl.r
    w(wsL, f"C{r}", dtxt(dk, "Trade payables ageing schedule as at "), fB); nl.r += 1
    r = nl.r
    for i, h in enumerate(["Particulars", "Less than 1 year", "1-2 years", "2-3 years", "More than 3 years"]):
        w(wsL, f"{'CDEFG'[i]}{r}", h, fB, bd=bTB, al=wrapc if i else None)
    nl.r += 1
    first = nl.r
    for j, lbl in enumerate(["(i) MSME", "(ii) Others", "(iii) Disputed dues - MSME", "(iv) Disputed dues - Others"]):
        r = nl.r
        w(wsL, f"C{r}", lbl)
        for i in range(4):
            w(wsL, f"{'DEFG'[i]}{r}", f"={q(CRA)}!{'CDEF'[i]}{base+j}/{SCALE}", fmt=ACC)
        nl.r += 1
    r = nl.r
    w(wsL, f"C{r}", "Total", fB, bd=bTot)
    for i in range(4):
        w(wsL, f"{'DEFG'[i]}{r}", f"=SUM({'DEFG'[i]}{first}:{'DEFG'[i]}{r-1})", fB, ACC, bd=bTot)
    w(wsL, f"I{r}", f"=ROUND(SUM(D{r}:G{r})-{'DEF'[k]}{NTOT['TP'][1]},2)", fS, ACC)
    nl.r += 2
nl.title("OCL", "Other current liabilities")
first = nl.r
for lbl in ["Statutory dues payable", "Salary and other employee dues payable", "Directors' remuneration payable", "Expenses and other payables",
            "Credit card dues", "Interest accrued on dues to MSME", "Advance from customers"]:
    nl.tb(lbl, "Cr")
nl.total(first, nl.r - 1, "OCL")
for i, c in enumerate("DEF"):
    w(wsL, f"{c}{msme_b}", f"={c}{NR['Interest accrued on dues to MSME'][1]}", fmt=ACC)
nl.gap()
nl.title("STP", "Short-term provisions")
a = nl.tb("Provision for income tax (net of advance tax)", "Cr")
nl.total(a, a, "STP")
NOTE_REF.update({"NOTE_" + k: nl.no[k] for k in ["LTB", "OLTL", "LTP", "STB", "TP", "OCL", "STP"]})

# ---------------- Notes_PPE
wsPP = sheet(NP, {"A": 1, "B": 5, "C": 40, "D": 16, "E": 16, "F": 16, "G": 16, "H": 3, "I": 12})
head(wsPP, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", span="G")
r = 8
w(wsPP, f"B{r}", f"={nl.no['STP']}+1", fB); w(wsPP, f"C{r}", "Property, Plant and Equipment and Intangible assets", fB)
NOTE_REF["NOTE_PPE"] = f"{q(NP)}!$B${r}"
r += 2
# summary (feeds BS_1)
w(wsPP, f"C{r}", "Net carrying amount", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24", "d23"]):
    w(wsPP, f"{'DEFG'[i]}{r}", dtxt(k), fB if i < 3 else fG, bd=bTB, al=wrapc)
wsPP.row_dimensions[r].height = 30
r += 1
SUMROW = {}
for nm in ["Property, Plant and Equipment", "Intangible assets", "Capital work-in-progress", "Intangible assets under development"]:
    SUMROW[nm] = r
    w(wsPP, f"C{r}", nm)
    r += 1
w(wsPP, f"C{r}", "Total", fB, bd=bTot)
for col in "DEFG": w(wsPP, f"{col}{r}", f"=SUM({col}{r-4}:{col}{r-1})", fB, ACC, bd=bTot)
r += 2
w(wsPP, f"C{r}", "The Company maintains its fixed assets on written down value basis in its books. The movement in the net carrying amount is given below.", fI)
r += 2
PPE_GROUPS = [("(a) Property, Plant and Equipment", ["Land", "Buildings", "Furniture and fittings", "Vehicles", "Office equipment", "Electrical installations", "Computers"]),
              ("(b) Intangible assets", ["Computer software"]), ("(c) Capital work-in-progress", ["Capital work-in-progress"]),
              ("(d) Intangible assets under development", ["Intangible assets under development"])]
PPE_TAB = {}
for fy in FYL:
    yk = END[fy]; pk = START[fy]
    w(wsPP, f"C{r}", dtxt(yk, "As at "), fB); r += 1
    w(wsPP, f"C{r}", "Particulars", fB, bd=bTB)
    w(wsPP, f"D{r}", f'="As at "&TEXT({md[pk]}+1,"mmmm dd, yyyy")', fB, bd=bTB, al=wrapc)
    w(wsPP, f"E{r}", "Additions during the year", fB, bd=bTB, al=wrapc)
    w(wsPP, f"F{r}", "Depreciation / deductions during the year", fB, bd=bTB, al=wrapc)
    w(wsPP, f"G{r}", dtxt(yk), fB, bd=bTB, al=wrapc)
    wsPP.row_dimensions[r].height = 30
    r += 1
    subs = []
    for gname, labels in PPE_GROUPS:
        w(wsPP, f"C{r}", gname, fB); r += 1
        first = r
        for lbl in labels:
            w(wsPP, f"C{r}", lbl)
            w(wsPP, f"D{r}", "=" + (tbf(PREVFY[fy], f"$C{r}") if fy in PREVFY else tb_open(f"$C{r}")), fmt=ACC)
            w(wsPP, f"E{r}", "=" + tbf(fy, f"$C{r}", "C") + f"/{SCALE}", fmt=ACC)
            w(wsPP, f"G{r}", "=" + tbf(fy, f"$C{r}"), fmt=ACC)
            w(wsPP, f"F{r}", f"=D{r}+E{r}-G{r}", fmt=ACC)
            NR.setdefault(lbl, (NP, r))
            r += 1
        w(wsPP, f"C{r}", "Total", fB, bd=bT)
        for col in "DEFG": w(wsPP, f"{col}{r}", f"=SUM({col}{first}:{col}{r-1})", fB, ACC, bd=bT)
        subs.append((gname, r))
        r += 1
    w(wsPP, f"C{r}", "Grand total", fB, bd=bTot)
    for col in "DEFG": w(wsPP, f"{col}{r}", "=" + "+".join(f"{col}{x}" for _, x in subs), fB, ACC, bd=bTot)
    PPE_TAB[fy] = (subs, r)
    r += 1
    w(wsPP, f"C{r}", "   Depreciation and amortisation charged to Statement of Profit and Loss", fI)
    w(wsPP, f"F{r}", "DEPLINK" + fy, fI, ACC); r += 1
    w(wsPP, f"C{r}", "   Deletions / transfers during the year", fI)
    w(wsPP, f"F{r}", f"=F{r-2}-F{r-1}", fI, ACC); r += 2
names4 = ["Property, Plant and Equipment", "Intangible assets", "Capital work-in-progress", "Intangible assets under development"]
for k, nm in enumerate(names4):
    rr = SUMROW[nm]
    for i, fy in enumerate(FYL):
        w(wsPP, f"{'DEF'[i]}{rr}", f"=G{PPE_TAB[fy][0][k][1]}", fmt=ACC)
    w(wsPP, f"G{rr}", f"=D{PPE_TAB['FY24'][0][k][1]}", fG, ACC)
    NTOT[nm] = (NP, rr)

# ---------------- Notes_Assets
wsNA = note_sheet(NA)
na = Note(wsNA, f"={NOTE_REF['NOTE_PPE']}+1")
na.title("DT", "Deferred tax assets / (liabilities) (net)")
a = na.tb("Deferred tax asset (net)", "Dr", adj=True)
b_ = na.tb("Deferred tax liability (net)", "Cr", adj=True)
NTOT["DTA"], NTOT["DTL"] = (NA, a), (NA, b_)
net_ = na.line("Net deferred tax asset / (liability)", lambda i, c: f"={c}{a}-{c}{b_}", bold=True, bd=bTot)
na.gap()
na.line("Movement during the year", bold=True)
mo = na.line("Opening balance", lambda i, c: f"={'EFG'[i]}{net_}", cols="DEF")
mc = na.line("Add: Credit / (charge) to the Statement of Profit and Loss", None)
mcl = na.line("Closing balance", lambda i, c: f"={c}{mo}+{c}{mc}", bold=True, bd=bT, cols="DEF")
for i, c in enumerate("DEF"): w(wsNA, f"{'IJK'[i]}{mcl}", f"=ROUND({c}{mcl}-{c}{net_},2)", fS, ACC)
na.text("Deferred tax arises mainly on account of difference between depreciation as per books and as per the Income-tax Act, and expenses allowable on payment basis (gratuity).", 2)
na.gap()
na.title("LTLA", "Long-term loans and advances")
na.line("Unsecured, considered good", bold=True)
a = na.tb("Security deposits", "Dr")
na.total(a, a, "LTLA")
na.gap()
na.title("INV", "Inventories")
na.line("(Valued at lower of cost and net realisable value, as taken, valued and certified by the management)", italic=True)
r = na.r
w(wsNA, f"C{r}", "Stock-in-trade")
w(wsNA, f"D{r}", f"={md['clstk']}/{SCALE}", fmt=ACC)
w(wsNA, f"E{r}", "=" + tbf("FY26", '"Opening stock"'), fmt=ACC)
w(wsNA, f"F{r}", "=" + tbf("FY25", '"Opening stock"'), fmt=ACC)
w(wsNA, f"G{r}", "=" + tbf("FY24", '"Opening stock"'), fG, ACC)
NR["Stock-in-trade"] = (NA, r)
na.r += 1
na.total(r, r, "INV")
na.text("Closing stock as at March 31, 2026 is as per the audited financial statements; closing stock of earlier years is the opening stock carried in the next year's books.", 1)
na.gap()
na.title("TR", "Trade receivables")
na.line("Unsecured, considered good", bold=True)
a = na.tb("Trade receivables", "Dr")
na.total(a, a, "TR")
na.gap()
for k, (dk, base) in enumerate([("d26", 5), ("d25", 13), ("d24", 21)]):
    r = na.r
    w(wsNA, f"C{r}", dtxt(dk, "Trade receivables ageing schedule as at "), fB); na.r += 1
    r = na.r
    hdrs = ["Particulars", "Less than 6 months", "6 months - 1 year", "1-2 years", "2-3 years", "More than 3 years"]
    wsNA.column_dimensions["H"].width = 14
    for i, h in enumerate(hdrs):
        w(wsNA, f"{'CDEFGH'[i]}{r}", h, fB, bd=bTB, al=wrapc if i else None)
    wsNA.row_dimensions[r].height = 26
    na.r += 1
    first = na.r
    for j, lbl in enumerate(["(i) Undisputed - considered good", "(ii) Undisputed - considered doubtful", "(iii) Disputed - considered good", "(iv) Disputed - considered doubtful"]):
        r = na.r
        w(wsNA, f"C{r}", lbl)
        for i in range(5):
            w(wsNA, f"{'DEFGH'[i]}{r}", f"={q(DRA)}!{'CDEFG'[i]}{base+j}/{SCALE}", fmt=ACC)
        na.r += 1
    r = na.r
    w(wsNA, f"C{r}", "Total", fB, bd=bTot)
    for i in range(5):
        w(wsNA, f"{'DEFGH'[i]}{r}", f"=SUM({'DEFGH'[i]}{first}:{'DEFGH'[i]}{r-1})", fB, ACC, bd=bTot)
    w(wsNA, f"I{r}", f"=ROUND(SUM(D{r}:H{r})-{'DEF'[k]}{NTOT['TR'][1]},2)", fS, ACC)
    na.r += 2
na.title("CB", "Cash and bank balances")
na.line("Cash and cash equivalents", bold=True)
a = na.tb("Cash on hand", "Dr")
b_ = na.tb("Balances with banks in current accounts", "Dr")
cce = na.line("Total cash and cash equivalents", lambda i, c: f"={c}{a}+{c}{b_}", bold=True, bd=bT)
NTOT["CCE"] = (NA, cce)
na.line("Other bank balances", bold=True)
fd = na.tb("Fixed deposits with banks", "Dr")
r = na.line("Total", lambda i, c: f"={c}{cce}+{c}{fd}", bold=True, bd=bTot)
NTOT["CB"] = (NA, r)
na.text("Fixed deposits include deposits held as margin money / lien against bank guarantees and overdraft facilities.", 1)
na.gap()
na.title("STLA", "Short-term loans and advances")
na.line("Unsecured, considered good", bold=True)
first = na.r
for lbl in ["Advance to suppliers", "Advances to employees", "Other advances"]:
    na.tb(lbl, "Dr")
na.total(first, na.r - 1, "STLA")
na.gap()
na.title("OCA", "Other current assets")
first = na.r
for lbl in ["Balances with government authorities", "Advance tax, TDS and TCS receivable", "Prepaid expenses", "Unbilled revenue", "Earnest money deposits",
            "Interest accrued on deposits", "IPO expenses pending adjustment", "Other current assets"]:
    na.tb(lbl, "Dr")
na.total(first, na.r - 1, "OCA")
NOTE_REF.update({"NOTE_" + k: na.no[k] for k in ["DT", "LTLA", "INV", "TR", "CB", "STLA", "OCA"]})

# ---------------- Notes_P&L
wsNP = note_sheet(NPL)
nP = Note(wsNP, f"={na.no['OCA']}+1")
nP.title("REV", "Revenue from operations", pl=True)
first = nP.r
nP.line("Sale of goods and services", bold=True)
for lbl in ["Sale of products", "Sale of services", "Other operating revenue"]: nP.tb(lbl, "Cr")
nP.total(first + 1, nP.r - 1, "REV")
nP.gap()
nP.title("OI", "Other income", pl=True)
a = nP.tb("Interest income", "Cr")
b_ = nP.tb("Other non-operating income", "Cr")
oi_t = nP.total(a, b_, "OI")
pbt_l = nP.line("Restated profit before tax", lambda i, c: f"={PL}!{c}{{PBT}}", italic=True)
nP.line("Other income as a percentage of restated profit before tax", lambda i, c: f"=IFERROR({c}{oi_t}/{c}{pbt_l},0)", fmt=PCT, italic=True)
nP.text("Interest income on fixed deposits and income tax refund is recurring in nature and not related to the main business activity. Other non-operating income "
        "(discount received, balances written back, exchange gain) is non-recurring in nature.", 2)
nP.gap()
nP.title("PUR", "Purchases of stock-in-trade", pl=True)
a = nP.tb("Purchases of stock-in-trade", "Dr")
b_ = nP.tb("Direct expenses", "Dr")
nP.total(a, b_, "PUR")
nP.gap()
nP.title("CHG", "Changes in inventories of stock-in-trade", pl=True)
a = nP.line("Inventories at the beginning of the year (a)", lambda i, c: "=" + nref("Stock-in-trade", "EFG"[i]))
b_ = nP.line("Inventories at the end of the year (b)", lambda i, c: "=" + nref("Stock-in-trade", "DEF"[i]))
r = nP.line("(Increase) / decrease in inventories (a - b)", lambda i, c: f"={c}{a}-{c}{b_}", bold=True, bd=bTot)
NTOT["CHG"] = (NPL, r)
nP.gap()
nP.title("EMP", "Employee benefits expense", pl=True)
first = nP.r
for lbl in ["Salaries, wages and incentives", "Directors' remuneration", "Contribution to provident and other funds"]: nP.tb(lbl, "Dr")
nP.tb("Gratuity", "Dr", adj=True)
nP.tb("Staff welfare expenses", "Dr")
nP.total(first, nP.r - 1, "EMP")
nP.gap()
nP.title("FIN", "Finance costs", pl=True)
first = nP.r
for lbl in ["Interest on borrowings", "Interest on delayed payments to MSME", "Other borrowing costs"]: nP.tb(lbl, "Dr")
nP.total(first, nP.r - 1, "FIN")
for i, c in enumerate("DEF"):
    w(wsL, f"{c}{msme_c}", "=" + nref("Interest on delayed payments to MSME", c), fmt=ACC)
nP.gap()
nP.title("DEP", "Depreciation and amortisation expense", pl=True)
a = nP.tb("Depreciation and amortisation", "Dr")
nP.total(a, a, "DEP")
nP.gap()
OE_LBL = ["Electricity charges", "Rent", "Rates and taxes", "Insurance", "Repairs and maintenance", "Travelling and conveyance", "Brokerage and commission",
          "Advertisement and business promotion", "Legal and professional charges", "Payment to auditors", "Communication, postage and courier",
          "Printing and stationery", "Freight and transportation", "Bad debts and balances written off", "Loss on foreign exchange fluctuation", "Miscellaneous expenses"]
nP.title("OE", "Other expenses", pl=True)
first = nP.r
for lbl in OE_LBL: nP.tb(lbl, "Dr")
nP.total(first, nP.r - 1, "OE")
nP.gap()
nP.line("Payment to auditors", bold=True)
nP.line("As statutory auditor (including tax audit)", lambda i, c: "=" + nref("Payment to auditors", c))
nP.line("For other services", 0)
nP.gap()
nP.title("TAX", "Tax expense", pl=True)
a = nP.tb("Current tax", "Dr")
b_ = nP.tb("Deferred tax", "Dr", adj=True)
nP.total(a, b_, "TAX")
nP.gap()
bookpat = nP.line("Net profit after tax as per books of account (memorandum)", lambda i, c: "=" + tbf(FYL[i], '"Net profit as per books"'), italic=True)
NR["__bookpat"] = (NPL, bookpat)
nP.gap()
nP.title("EPS", "Earnings per share", pl=True)
EPS_START = nP.r
NOTE_REF.update({"NOTE_" + k: nP.no[k] for k in ["REV", "OI", "PUR", "CHG", "EMP", "FIN", "DEP", "OE", "TAX", "EPS"]})
for i, c in enumerate("DEF"):
    w(wsNA, f"{c}{mc}", f"=-{nref('Deferred tax', c)}", fmt=ACC)
for fy in FYL:
    for row in wsPP.iter_rows(min_col=6, max_col=6):
        for c in row:
            if c.value == "DEPLINK" + fy:
                c.value = "=" + tref("DEP", "DEF"[FYL.index(fy)])

# ================================================================== BS_1 (totals linked from notes)
wsB = sheet(BS, {"A": 1, "B": 52, "C": 11, "D": 16.8, "E": 16.8, "F": 16.8, "G": 3, "H": 15})
head(wsB, "ANNEXURE I - RESTATED SUMMARY STATEMENT OF ASSETS AND LIABILITIES")
w(wsB, "B8", "Particulars", fB, bd=bTB, al=Alignment(vertical="center"))
w(wsB, "C8", "Annexure VI Notes", fB, bd=bTB, al=wrapc)
for col, k in zip("DEF", ["d26", "d25", "d24"]):
    w(wsB, f"{col}8", dtxt(k), fB, bd=bTB, al=wrapc)
w(wsB, "H8", dtxt("d23"), fG, bd=bTB, al=wrapc)
wsB.row_dimensions[8].height = 30
FACE = {}
r = 9


def face(label, key=None, bold=False, note=None, indent=0):
    global r
    w(wsB, f"B{r}", ("   " * indent) + label, fB if bold else fN)
    if note: w(wsB, f"C{r}", "=" + NOTE_REF[note], al=Alignment(horizontal="center"))
    if key:
        for col, nc in zip("DEFH", "DEFG"):
            w(wsB, f"{col}{r}", "=" + tref(key, nc), fG if col == "H" else fN, ACC)
        FACE[key] = r
    r += 1
    return r - 1


def tot(label, rows_, key, dbl_=False):
    global r
    w(wsB, f"B{r}", label, fB)
    for col in "DEFH":
        w(wsB, f"{col}{r}", "=SUM(" + ",".join(f"{col}{x}" for x in rows_) + ")" if len(rows_) > 2 else "=" + "+".join(f"{col}{x}" for x in rows_),
          fG if col == "H" else fB, ACC, bd=bTot if dbl_ else bT)
    FACE[key] = r
    r += 1
    return r - 1


face("EQUITY AND LIABILITIES", bold=True); r += 1
face("Shareholders' funds", bold=True)
a = face("Share capital", "SC", note="NOTE_SC")
b = face("Reserves and surplus", "RS", note="NOTE_RS")
e1 = tot("Total Shareholders' funds", [a, b], "EQ"); r += 1
face("Non-current liabilities", bold=True)
n1 = face("Long-term borrowings", "LTB", note="NOTE_LTB")
n2 = face("Deferred tax liabilities (Net)", "DTL", note="NOTE_DT")
n3 = face("Other long-term liabilities", "OLTL", note="NOTE_OLTL")
n4 = face("Long-term provisions", "LTP", note="NOTE_LTP")
e2 = tot("Total Non-current liabilities", [n1, n2, n3, n4], "NCL"); r += 1
face("Current liabilities", bold=True)
c1 = face("Short-term borrowings", "STB", note="NOTE_STB")
face("Trade payables", note="NOTE_TP")
c2 = face("- Total outstanding dues of micro and small enterprises", "TP_MSME")
c3 = face("- Total outstanding dues of creditors other than micro and small enterprises", "TP_OTH")
c4 = face("Other current liabilities", "OCL", note="NOTE_OCL")
c5 = face("Short-term provisions", "STP", note="NOTE_STP")
e3 = tot("Total Current liabilities", [c1, c2, c3, c4, c5], "CL"); r += 1
TL = tot("Total Equity and Liabilities", [e1, e2, e3], "TL", True); r += 1
face("ASSETS", bold=True); r += 1
face("Non-current assets", bold=True)
face("Property, Plant and Equipment and Intangible assets", note="NOTE_PPE")
a1 = face("Property, Plant and Equipment", "Property, Plant and Equipment", indent=1)
a2 = face("Intangible assets", "Intangible assets", indent=1)
a3 = face("Capital work-in-progress", "Capital work-in-progress", indent=1)
a4 = face("Intangible assets under development", "Intangible assets under development", indent=1)
a5 = face("Deferred tax assets (Net)", "DTA", note="NOTE_DT")
a6 = face("Long-term loans and advances", "LTLA", note="NOTE_LTLA")
f1 = tot("Total Non-current assets", [a1, a2, a3, a4, a5, a6], "NCA"); r += 1
face("Current assets", bold=True)
b1 = face("Inventories", "INV", note="NOTE_INV")
b2 = face("Trade receivables", "TR", note="NOTE_TR")
b3 = face("Cash and bank balances", "CB", note="NOTE_CB")
b4 = face("Short-term loans and advances", "STLA", note="NOTE_STLA")
b5 = face("Other current assets", "OCA", note="NOTE_OCA")
f2 = tot("Total Current assets", [b1, b2, b3, b4, b5], "CA"); r += 1
TA = tot("Total Assets", [f1, f2], "TA", True)
for col in "DEFH":
    w(wsB, f"{col}1", f"=ROUND({col}{TL}-{col}{TA},2)", fS, ACC)
r += 1
NOTE_TEXT_ROW = r
w(wsB, f"B{r}", "The above statement should be read with Annexure IV - Background, basis of preparation and significant accounting policies, Annexure V - "
                "Statement of adjustments to audited financial statements and Annexure VI - Notes to the restated summary statements.", fN, al=just)
wsB.merge_cells(f"B{r}:F{r+1}"); r += 3
sign_block(wsB, r)

# ================================================================== PL_2 (totals linked from notes)
wsP = sheet(PL, {"A": 1, "B": 60, "C": 11, "D": 16.8, "E": 16.8, "F": 16.8})
head(wsP, "ANNEXURE II - RESTATED SUMMARY STATEMENT OF PROFIT AND LOSS")
w(wsP, "B7", "Particulars", fB, bd=bTB, al=Alignment(vertical="center"))
w(wsP, "C7", "Annexure VI Notes", fB, bd=bTB, al=wrapc)
for col, k in zip("DEF", ["d26", "d25", "d24"]):
    w(wsP, f"{col}7", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
wsP.row_dimensions[7].height = 30
PF = {}
r = 8


def pface(label, ref=None, note=None, bold=False, formula=None, bd=None, key=None, fmt=ACC):
    global r
    w(wsP, f"B{r}", label, fB if bold else fN, bd=bd)
    if note: w(wsP, f"C{r}", "=" + NOTE_REF[note], al=Alignment(horizontal="center"))
    for col in "DEF":
        v = ("=" + ref(col)) if ref else (formula(col) if formula else None)
        if v is not None: w(wsP, f"{col}{r}", v, fB if bold else fN, fmt, bd=bd)
    if key: PF[key] = r
    r += 1
    return r - 1


pface("Income", bold=True)
i1 = pface("I. Revenue from operations", lambda c: tref("REV", c), "NOTE_REV", key="REV")
i2 = pface("II. Other income", lambda c: tref("OI", c), "NOTE_OI", key="OI")
i3 = pface("III. Total Income (I + II)", formula=lambda c: f"=SUM({c}{i1}:{c}{i2})", bold=True, bd=bT, key="TI")
pface("IV. Expenses", bold=True)
x1 = pface("Purchases of stock-in-trade", lambda c: tref("PUR", c), "NOTE_PUR", key="PUR")
x2 = pface("Changes in inventories of stock-in-trade", lambda c: tref("CHG", c), "NOTE_CHG", key="CHG")
x3 = pface("Employee benefits expense", lambda c: tref("EMP", c), "NOTE_EMP", key="EMP")
x4 = pface("Finance costs", lambda c: tref("FIN", c), "NOTE_FIN", key="FIN")
x5 = pface("Depreciation and amortisation expense", lambda c: tref("DEP", c), "NOTE_DEP", key="DEP")
x6 = pface("Other expenses", lambda c: tref("OE", c), "NOTE_OE", key="OE")
x7 = pface("Total Expenses", formula=lambda c: f"=SUM({c}{x1}:{c}{x6})", bold=True, bd=bT, key="TE")
v1 = pface("V. Profit before exceptional and extraordinary items and tax (III - IV)", formula=lambda c: f"={c}{i3}-{c}{x7}", bold=True)
v2 = pface("VI. Exceptional and extraordinary items", formula=lambda c: "=0")
pbt = pface("VII. Profit before tax (V - VI)", formula=lambda c: f"={c}{v1}+{c}{v2}", bold=True, bd=bT, key="PBT")
pface("VIII. Tax expense:", bold=True)
t1 = pface("Current tax", lambda c: nref("Current tax", c), "NOTE_TAX", key="TAXC")
t2 = pface("Deferred tax", lambda c: nref("Deferred tax", c), key="TAXD")
t3 = pface("Total tax expense", formula=lambda c: f"=SUM({c}{t1}:{c}{t2})", bd=bT, key="TAX")
pat = pface("IX. Profit for the year (VII - VIII)", formula=lambda c: f"={c}{pbt}-{c}{t3}", bold=True, bd=bTot, key="PAT")
r += 1
pface("X. Earnings per equity share (Face value of Rs. 10 each)", bold=True)
eps_r = pface("Basic and Diluted (in Rs.)", note="NOTE_EPS", key="EPS")
r += 1
w(wsP, f"B{r}", f"={BS}!B{NOTE_TEXT_ROW}", fN, al=just); wsP.merge_cells(f"B{r}:F{r+1}"); r += 3
sign_block(wsP, r)
# check: restated PAT = PAT as per books + restatement adjustments
for i, c in enumerate("DEF"):
    fy = FYL[i]
    crit = f'{AR("D")},"P&L",{AR("A")},">"&{md[START[fy]]},{AR("A")},"<="&{md[END[fy]]}'
    w(wsP, f"{c}1", f"=ROUND({c}{pat}-{nref('__bookpat', c)}+(SUMIFS({AR('E')},{crit})-SUMIFS({AR('F')},{crit})),2)", fS, ACC)
# resolve {PAT}/{PBT} placeholders in notes
for ws_ in (wsE, wsNP):
    for row in ws_.iter_rows():
        for c in row:
            if isinstance(c.value, str) and ("{PAT}" in c.value or "{PBT}" in c.value):
                c.value = c.value.replace("{PAT}", str(pat)).replace("{PBT}", str(pbt))

# ================================================================== CFS_3 (from BS_1, PL_2 and notes)
wsC = sheet(CF, {"A": 1, "B": 62, "C": 18, "D": 18, "E": 18})
head(wsC, "ANNEXURE III - RESTATED SUMMARY STATEMENT OF CASH FLOWS", span="E")
w(wsC, "B8", "Particulars", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsC, f"{'CDE'[i]}8", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
wsC.row_dimensions[8].height = 30
CUR, PRV = "DEF", "EFH"
NCUR, NPRV = "DEF", "EFG"
r = 9


def cf(label, fn=None, bold=False, bd=None):
    global r
    w(wsC, f"B{r}", label, fB if bold else fN, bd=bd)
    if fn:
        for i, col in enumerate("CDE"):
            w(wsC, f"{col}{r}", fn(i, CUR[i], PRV[i]), fB if bold else fN, ACC, bd=bd)
    r += 1
    return r - 1


def Bf(key, c): return f"{BS}!{c}{FACE[key]}"
def P(key, i): return f"{PL}!{'DEF'[i]}{PF[key]}"
def Nn(lbl, i, prev=False): return nref(lbl, (NPRV if prev else NCUR)[i])


cf("A) Cash flows from operating activities", bold=True)
a0 = cf("Net profit before tax", lambda i, c, p: "=" + P("PBT", i))
cf("Adjustments for:")
a1_ = cf("Depreciation and amortisation expense", lambda i, c, p: "=" + P("DEP", i))
a2_ = cf("Finance costs", lambda i, c, p: "=" + P("FIN", i))
a3_ = cf("Interest income", lambda i, c, p: "=-" + Nn("Interest income", i))
a4_ = cf("Adjustments made directly in reserves", lambda i, c, p: f"={q(NE)}!{NCUR[i]}{RSR['book']}+{q(NE)}!{NCUR[i]}{RSR['direct']}")
a5_ = cf("Operating profit before working capital changes", lambda i, c, p: f"=SUM({'CDE'[i]}{a0}:{'CDE'[i]}{a4_})", bold=True, bd=bT)
cf("Adjustments for changes in working capital:")
cf("(Increase) / decrease in inventories", lambda i, c, p: f"={Bf('INV', p)}-{Bf('INV', c)}")
cf("(Increase) / decrease in trade receivables", lambda i, c, p: f"={Bf('TR', p)}-{Bf('TR', c)}")
cf("(Increase) / decrease in short-term loans and advances", lambda i, c, p: f"={Bf('STLA', p)}-{Bf('STLA', c)}")
cf("(Increase) / decrease in long-term loans and advances", lambda i, c, p: f"={Bf('LTLA', p)}-{Bf('LTLA', c)}")
cf("(Increase) / decrease in other current assets",
   lambda i, c, p: f"=({Bf('OCA', p)}-{Nn('Advance tax, TDS and TCS receivable', i, True)}-{Nn('Interest accrued on deposits', i, True)})"
                   f"-({Bf('OCA', c)}-{Nn('Advance tax, TDS and TCS receivable', i)}-{Nn('Interest accrued on deposits', i)})")
cf("Increase / (decrease) in trade payables", lambda i, c, p: f"=({Bf('TP_MSME', c)}+{Bf('TP_OTH', c)})-({Bf('TP_MSME', p)}+{Bf('TP_OTH', p)})")
cf("Increase / (decrease) in other current liabilities", lambda i, c, p: f"={Bf('OCL', c)}-{Bf('OCL', p)}")
w8 = cf("Increase / (decrease) in long-term provisions and other long-term liabilities",
        lambda i, c, p: f"=({Bf('LTP', c)}+{Bf('OLTL', c)})-({Bf('LTP', p)}+{Bf('OLTL', p)})")
w9 = cf("Cash generated from operations", lambda i, c, p: f"=SUM({'CDE'[i]}{a5_}:{'CDE'[i]}{w8})", bold=True, bd=bT)
w10 = cf("Income taxes paid (net of refunds)",
         lambda i, c, p: f"=-{P('TAX', i)}+({Bf('STP', c)}-{Bf('STP', p)})-({Nn('Advance tax, TDS and TCS receivable', i)}-{Nn('Advance tax, TDS and TCS receivable', i, True)})"
                         f"-({Bf('DTA', c)}-{Bf('DTA', p)})+({Bf('DTL', c)}-{Bf('DTL', p)})")
A_ = cf("Net cash flow from / (used in) operating activities (A)", lambda i, c, p: f"={'CDE'[i]}{w9}+{'CDE'[i]}{w10}", bold=True, bd=bTot)
r += 1
cf("B) Cash flows from investing activities", bold=True)
fa = names4
i1_ = cf("Purchase of Property, Plant and Equipment and intangible assets (net of disposals)",
         lambda i, c, p: "=-((" + "+".join(Bf(x, c) for x in fa) + ")-(" + "+".join(Bf(x, p) for x in fa) + f")+{P('DEP', i)})")
cf("(Investment in) / redemption of fixed deposits with banks", lambda i, c, p: f"={Nn('Fixed deposits with banks', i, True)}-{Nn('Fixed deposits with banks', i)}")
i3_ = cf("Interest received", lambda i, c, p: f"={Nn('Interest income', i)}-({Nn('Interest accrued on deposits', i)}-{Nn('Interest accrued on deposits', i, True)})")
B_ = cf("Net cash flow from / (used in) investing activities (B)", lambda i, c, p: f"=SUM({'CDE'[i]}{i1_}:{'CDE'[i]}{i3_})", bold=True, bd=bTot)
r += 1
cf("C) Cash flows from financing activities", bold=True)
f1_ = cf("Proceeds from issue of share capital", lambda i, c, p: f"={Bf('SC', c)}-{Bf('SC', p)}")
cf("Proceeds from / (repayment of) long-term borrowings (net)", lambda i, c, p: f"={Bf('LTB', c)}-{Bf('LTB', p)}")
cf("Proceeds from / (repayment of) short-term borrowings (net)", lambda i, c, p: f"={Bf('STB', c)}-{Bf('STB', p)}")
f4_ = cf("Finance costs paid", lambda i, c, p: f"=-{P('FIN', i)}")
C_ = cf("Net cash flow from / (used in) financing activities (C)", lambda i, c, p: f"=SUM({'CDE'[i]}{f1_}:{'CDE'[i]}{f4_})", bold=True, bd=bTot)
r += 1
nt = cf("Net increase / (decrease) in cash and cash equivalents (A+B+C)", lambda i, c, p: f"={'CDE'[i]}{A_}+{'CDE'[i]}{B_}+{'CDE'[i]}{C_}", bold=True)
op = cf("Cash and cash equivalents at the beginning of the year", lambda i, c, p: "=" + tref("CCE", NPRV[i]))
clo = cf("Cash and cash equivalents at the end of the year", lambda i, c, p: f"={'CDE'[i]}{nt}+{'CDE'[i]}{op}", bold=True, bd=bTot)
r += 1
cf("Components of cash and cash equivalents:", bold=True)
k1 = cf("Cash on hand", lambda i, c, p: "=" + Nn("Cash on hand", i))
k2 = cf("Balances with banks in current accounts", lambda i, c, p: "=" + Nn("Balances with banks in current accounts", i))
k3 = cf("Total cash and cash equivalents", lambda i, c, p: f"=SUM({'CDE'[i]}{k1}:{'CDE'[i]}{k2})", bold=True, bd=bTot)
for i, col in enumerate("CDE"):
    w(wsC, f"{col}1", f"=ROUND({col}{k3}-{col}{clo},2)", fS, ACC)
r += 1
w(wsC, f"B{r}", "Note: The Cash Flow Statement has been prepared under the indirect method as set out in Accounting Standard 3 'Cash Flow Statements'. Bank overdrafts "
                "are shown under short-term borrowings and their movement is included in financing activities.", al=just)
wsC.merge_cells(f"B{r}:E{r+1}"); r += 3
sign_block(wsC, r, "B", "C", "E")

# ================================================================== RPT (from notes)
wsT = sheet(RPT, {"A": 1, "B": 5, "C": 52, "D": 16.8, "E": 16.8, "F": 16.8})
head(wsT, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS")
r = 8
w(wsT, f"B{r}", f"={nP.no['EPS']}+1", fB); w(wsT, f"C{r}", "Related party transactions", fB)
r += 2
w(wsT, f"C{r}", "(a) List of related parties and nature of relationship", fB); r += 1
w(wsT, f"C{r}", "Name of the party", fB, bd=bTB); w(wsT, f"D{r}", "Nature of relationship", fB, bd=bTB); wsT.merge_cells(f"D{r}:F{r}"); r += 1
w(wsT, f"C{r}", "Key Managerial Personnel", fBI); r += 1
for nm, rel in [(f"=PROPER({md['dir1']})", "Director and Promoter"), (f"=PROPER({md['dir2']})", "Director and Promoter"), (f"=PROPER({md['dir3']})", "Director")]:
    w(wsT, f"C{r}", nm); w(wsT, f"D{r}", rel); r += 1
w(wsT, f"C{r}", "Relatives of KMP / entities in which KMP have significant influence", fBI); r += 1
w(wsT, f"C{r}", "[To be furnished by the management]", fI); r += 2
w(wsT, f"C{r}", "(b) Transactions with related parties during the year", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsT, f"{'DEF'[i]}{r}", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
wsT.row_dimensions[r].height = 30
r += 1


def rp(label, fn, bold=False):
    global r
    w(wsT, f"C{r}", label, fB if bold else fN)
    if fn:
        for i, col in enumerate("DEF"):
            w(wsT, f"{col}{r}", fn(i, col), fB if bold else fN, ACC)
    r += 1


rp("Directors' remuneration (aggregate)", lambda i, c: "=" + nref("Directors' remuneration", c))
for d_ in DIRS:
    nm = d_.split(" - ")[1]
    rp(f"Loan accepted - {nm}", lambda i, c, d_=d_: f"={q(NL)}!{c}{DIRMOV[('acc', d_)]}")
    rp(f"Loan repaid - {nm}", lambda i, c, d_=d_: f"={q(NL)}!{c}{DIRMOV[('rep', d_)]}")
r += 1
w(wsT, f"C{r}", "(c) Outstanding balances at the end of the year", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsT, f"{'DEF'[i]}{r}", dtxt(k), fB, bd=bTB, al=wrapc)
wsT.row_dimensions[r].height = 30
r += 1
for d_ in DIRS:
    rp("Unsecured loan payable - " + d_.split(" - ")[1], lambda i, c, d_=d_: "=" + nref(d_, c))
rp("Directors' remuneration payable", lambda i, c: "=" + nref("Directors' remuneration payable", c))
r += 1
w(wsT, f"C{r}", "Note: Related party relationships are as identified by the management. Director-wise split of remuneration and other transactions (rent, reimbursements, "
                "guarantees) are to be furnished by the management.", al=just)
wsT.merge_cells(f"C{r}:F{r+1}")
RPT_NO = f"{q(RPT)}!$B$8"

# ================================================================== Gratuity and contingencies
wsG = sheet(GC, {"A": 1, "B": 5, "C": 66, "D": 16.8, "E": 16.8, "F": 16.8})
head(wsG, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS")
r = 8
w(wsG, f"B{r}", f"={RPT_NO}+1", fB); w(wsG, f"C{r}", "Contingent liabilities and commitments", fB)
r += 2
w(wsG, f"C{r}", "a) Contingent liabilities", fB); r += 1
w(wsG, f"C{r}", "Particulars", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsG, f"{'DEF'[i]}{r}", dtxt(k), fB, bd=bTB, al=wrapc)
wsG.row_dimensions[r].height = 30
r += 1
first = r
for lbl in ["(a) Claims against the Company not acknowledged as debt", "(b) Guarantees issued by banks on behalf of the Company", "(c) Other money for which the Company is contingently liable"]:
    w(wsG, f"C{r}", lbl)
    for col in "DEF": w(wsG, f"{col}{r}", 0, fBlue, ACC)
    r += 1
w(wsG, f"C{r}", "Total", fB, bd=bTot)
for col in "DEF": w(wsG, f"{col}{r}", f"=SUM({col}{first}:{col}{r-1})", fB, ACC, bd=bTot)
r += 2
w(wsG, f"C{r}", "b) Commitments", fB); r += 1
for lbl in ["Estimated amount of contracts remaining to be executed on capital account and not provided for", "Other commitments"]:
    w(wsG, f"C{r}", lbl)
    for col in "DEF": w(wsG, f"{col}{r}", "Nil", al=Alignment(horizontal="right"))
    r += 1
r += 1
EB_ROW = r
w(wsG, f"B{r}", "=B8+1", fB); w(wsG, f"C{r}", "Employee benefit plans", fB); r += 2
w(wsG, f"C{r}", "A. Defined contribution plan", fB); r += 1
w(wsG, f"C{r}", "Particulars", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsG, f"{'DEF'[i]}{r}", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
wsG.row_dimensions[r].height = 30
r += 1
w(wsG, f"C{r}", "Employer's contribution to provident fund and other funds")
for col in "DEF": w(wsG, f"{col}{r}", "=" + nref("Contribution to provident and other funds", col), fmt=ACC)
r += 2
w(wsG, f"C{r}", "B. Defined benefit obligation - Gratuity (unfunded)", fB); r += 1
w(wsG, f"C{r}", "The Company provides for gratuity to eligible employees as per the Payment of Gratuity Act, 1972. The scheme is unfunded.", al=just); wsG.merge_cells(f"C{r}:F{r}"); r += 1
w(wsG, f"C{r}", "Particulars", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsG, f"{'DEF'[i]}{r}", dtxt(k), fB, bd=bTB, al=wrapc)
wsG.row_dimensions[r].height = 30
r += 1
w(wsG, f"C{r}", "Gratuity expense recognised in the Statement of Profit and Loss")
for col in "DEF": w(wsG, f"{col}{r}", "=" + nref("Gratuity", col), fmt=ACC)
r += 1
w(wsG, f"C{r}", "Present value of obligation at the end of the year (Provision for gratuity)")
for col in "DEF": w(wsG, f"{col}{r}", "=" + nref("Provision for gratuity", col), fmt=ACC)
r += 1
w(wsG, f"C{r}", "Actuarial assumptions (discount rate, salary escalation, attrition, mortality) - as per actuarial valuation report, to be incorporated.", fI)

# ================================================================== Tax Computation (statement of tax shelter)
wsX = sheet(TX, {"A": 1, "B": 5, "C": 66, "D": 16.8, "E": 16.8, "F": 16.8})
head(wsX, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS")
r = 8
w(wsX, f"B{r}", f"={q(GC)}!B{EB_ROW}+1", fB); w(wsX, f"C{r}", "Statement of tax shelter", fB)
r += 2
w(wsX, f"C{r}", "Particulars", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsX, f"{'DEF'[i]}{r}", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
wsX.row_dimensions[r].height = 30
r += 1
TXR = {}


def tx(label, fn, key=None, fmt=ACC, bold=False, inp=False):
    global r
    w(wsX, f"C{r}", label, fB if bold else fN)
    for i, col in enumerate("DEF"):
        v = fn(i, col) if callable(fn) else fn
        if v is not None: w(wsX, f"{col}{r}", v, fBlue if inp else (fB if bold else fN), fmt, fill=yel if inp else None)
    if key: TXR[key] = r
    r += 1
    return r - 1


tx("Restated profit before tax (A)", lambda i, c: "=" + P("PBT", i), "pbt", bold=True)
tx("Normal tax rate (Section 115BAA)", lambda i, c: f"={md['tax']}", "rate", "0.000%")
tx("Notional tax at normal rate (B)", lambda i, c: f"={c}{TXR['pbt']}*{c}{TXR['rate']}", "nt")
r += 1
tx("Permanent differences", None, bold=True)
p1 = tx("Disallowance u/s 37 (penalties, interest on TDS and other non-deductible expenses)", 0, inp=True)
p2 = tx("Interest on delayed payments to MSME (Sec 23 of MSMED Act)", lambda i, c: "=" + nref("Interest on delayed payments to MSME", c))
pt = tx("Total permanent differences (C)", lambda i, c: f"={c}{p1}+{c}{p2}", bold=True)
r += 1
tx("Timing differences", None, bold=True)
t1_ = tx("Depreciation as per Companies Act", lambda i, c: "=" + P("DEP", i))
t2_ = tx("Less: Depreciation as per Income-tax Act", 0, inp=True)
t3_ = tx("Provision for gratuity disallowed u/s 43B", lambda i, c: "=" + nref("Gratuity", c))
t4_ = tx("Less: Gratuity paid during the year", 0, inp=True)
t5_ = tx("Other timing differences", 0, inp=True)
tt = tx("Total timing differences (D)", lambda i, c: f"={c}{t1_}-{c}{t2_}+{c}{t3_}-{c}{t4_}+{c}{t5_}", bold=True)
r += 1
ti = tx("Net adjustments (E = C + D)", lambda i, c: f"={c}{pt}+{c}{tt}", bold=True)
tx("Taxable income (A + E)", lambda i, c: f"={c}{TXR['pbt']}+{c}{ti}", "ti", bold=True)
tx("Tax payable at normal rate", lambda i, c: f"=MAX(0,{c}{TXR['ti']})*{c}{TXR['rate']}", "tp")
tx("Current tax as per Restated Summary Statement of Profit and Loss", lambda i, c: "=" + P("TAXC", i), "ct")
tx("Difference", lambda i, c: f"={c}{TXR['tp']}-{c}{TXR['ct']}")
r += 1
w(wsX, f"C{r}", "Note: Depreciation as per the Income-tax Act and payments under Sec 43B are to be taken from the computation of total income / tax audit report of the respective years.", fI, al=just)
wsX.merge_cells(f"C{r}:F{r+1}")

# ================================================================== Accounting ratios
wsR = sheet(AR_, {"A": 1, "B": 5, "C": 70, "D": 16.8, "E": 16.8, "F": 16.8})
head(wsR, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS")
r = 8
w(wsR, f"B{r}", f"={q(TX)}!B8+1", fB); w(wsR, f"C{r}", "Summary of accounting ratios", fB)
r += 2
w(wsR, f"C{r}", "Particulars", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsR, f"{'DEF'[i]}{r}", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
wsR.row_dimensions[r].height = 30
r += 1
RT = {}


def rl(label, fn, key=None, fmt=ACC, bold=False):
    global r
    w(wsR, f"C{r}", label, fB if bold else fN)
    for i, col in enumerate("DEF"):
        w(wsR, f"{col}{r}", fn(i, col), fB if bold else fN, fmt)
    if key: RT[key] = r
    r += 1
    return r - 1


rl("Total revenue from operations (A)", lambda i, c: "=" + P("REV", i), "rev")
rl("Net profit as restated (B)", lambda i, c: "=" + P("PAT", i), "pat")
rl("Add: Tax expense", lambda i, c: "=" + P("TAX", i), "tax")
rl("Add: Depreciation and amortisation", lambda i, c: "=" + P("DEP", i), "dep")
rl("Add: Finance costs", lambda i, c: "=" + P("FIN", i), "fin")
rl("Less: Other income", lambda i, c: "=-" + P("OI", i), "oi")
rl("EBITDA (C)", lambda i, c: f"=SUM({c}{RT['pat']}:{c}{RT['oi']})", "ebitda", bold=True)
rl("EBITDA margin (%) (C / A)", lambda i, c: f"=IFERROR({c}{RT['ebitda']}/{c}{RT['rev']},0)", fmt=PCT)
r += 1
rl("Net worth as restated (D)", lambda i, c: "=" + Bf("EQ", c), "nw")
rl("Return on net worth (%) (B / D)", lambda i, c: f"=IFERROR({c}{RT['pat']}/{c}{RT['nw']},0)", "ronw", PCT, True)
r += 1
rl("Number of equity shares outstanding at the end of the year (E)", lambda i, c: "=" + nref("__shares", c), "sh", ACC0)
rl("Weighted average number of equity shares (F)", lambda i, c: f"=({c}{RT['sh']}+{q(NE)}!{c}{sh_b})/2", "wa", ACC0)
rl("Weighted average number of equity shares adjusted for bonus / split (G)", lambda i, c: f"={c}{RT['wa']}*{md['bonus']}", "wab", ACC0)
rl("Basic and diluted earnings per share (Rs.) (B / G)", lambda i, c: f"=IFERROR({c}{RT['pat']}*{SCALE}/{c}{RT['wab']},0)", "eps", "#,##0.00", True)
rl("Net asset value per equity share (Rs.) (D / E adjusted for bonus / split)", lambda i, c: f"=IFERROR({c}{RT['nw']}*{SCALE}/({c}{RT['sh']}*{md['bonus']}),0)", "nav", "#,##0.00", True)
rl("Current assets (H)", lambda i, c: "=" + Bf("CA", c), "ca")
rl("Current liabilities (I)", lambda i, c: "=" + Bf("CL", c), "cl")
rl("Current ratio (H / I)", lambda i, c: f"=IFERROR({c}{RT['ca']}/{c}{RT['cl']},0)", "cr", '0.00', True)
r += 1
for t in ["Notes:", "1. Ratios have been calculated as below:",
          "Earnings per share (Rs.) = Restated profit after tax / Weighted average number of equity shares outstanding during the year (adjusted for bonus / split)",
          "Return on net worth (%) = Restated profit after tax / Restated net worth at the end of the year",
          "Net asset value per equity share (Rs.) = Restated net worth at the end of the year / Number of equity shares outstanding at the end of the year",
          "EBITDA = Restated profit after tax + tax expense + depreciation and amortisation + finance costs - other income",
          "2. \"Net worth\" means the aggregate value of the paid-up share capital and reserves and surplus as per the Restated Summary Statement of Assets and Liabilities.",
          "3. The figures disclosed above are based on the Restated Summary Statements of the Company."]:
    w(wsR, f"C{r}", t, fB if t == "Notes:" else fN, al=just); wsR.merge_cells(f"C{r}:F{r}"); r += 1
nP.r = EPS_START
nP.line("Restated profit for the year attributable to equity shareholders", lambda i, c: "=" + P("PAT", i))
nP.line("Weighted average number of equity shares (adjusted for bonus / split)", lambda i, c: f"={q(AR_)}!{c}{RT['wab']}", fmt=ACC0)
eps_note = nP.line("Basic and diluted earnings per share (Rs.)", lambda i, c: f"={q(AR_)}!{c}{RT['eps']}", fmt="#,##0.00", bold=True)
nP.line("Face value per equity share (Rs.)", lambda i, c: f"={md['fv']}", fmt="#,##0")
for col in "DEF":
    w(wsP, f"{col}{eps_r}", f"={q(NPL)}!{col}{eps_note}", fN, "#,##0.00")

# ================================================================== Analytical Ratios
wsY = sheet(AN, {"A": 1, "B": 5, "C": 34, "D": 30, "E": 30, "F": 13, "G": 13, "H": 13, "I": 12, "J": 12, "K": 40})
head(wsY, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", span="K")
r = 8
w(wsY, f"B{r}", f"={q(AR_)}!B8+1", fB); w(wsY, f"C{r}", "Analytical ratios", fB)
r += 2
for i, h in enumerate(["Ratio", "Numerator", "Denominator", None, None, None, "% Variance 2026 vs 2025", "% Variance 2025 vs 2024", "Reason for variance (more than 25%)"]):
    if h: w(wsY, f"{'CDEFGHIJK'[i]}{r}", h, fB, bd=bTB, al=wrapc)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsY, f"{'FGH'[i]}{r}", f'=TEXT({md[k]},"mmmm dd, yyyy")', fB, bd=bTB, al=wrapc)
wsY.row_dimensions[r].height = 30
r += 1
avg = lambda key, i: f"(({Bf(key, CUR[i])})+({Bf(key, PRV[i])}))/2"
debt = lambda cc: f"({Bf('LTB', cc)}+{Bf('STB', cc)})"
ratios = [
    ("Current ratio", "Current assets", "Current liabilities", lambda i: f"=IFERROR({Bf('CA', CUR[i])}/{Bf('CL', CUR[i])},0)", '0.00'),
    ("Debt-equity ratio", "Total debt", "Shareholder's equity", lambda i: f"=IFERROR({debt(CUR[i])}/{Bf('EQ', CUR[i])},0)", '0.00'),
    ("Debt service coverage ratio", "Earnings available for debt service", "Debt service",
     lambda i: f"=IFERROR(({P('PAT', i)}+{P('DEP', i)}+{P('FIN', i)})/({P('FIN', i)}+{'FGH'[i]}{{REPAY}}),0)", '0.00'),
    ("Return on equity ratio", "Net profit after taxes", "Average shareholders' equity", lambda i: f"=IFERROR({P('PAT', i)}/{avg('EQ', i)},0)", PCT),
    ("Inventory turnover ratio", "Cost of goods sold", "Average inventory", lambda i: f"=IFERROR(({P('PUR', i)}+{P('CHG', i)})/{avg('INV', i)},0)", '0.00'),
    ("Trade receivables turnover ratio", "Net credit sales", "Average trade receivables", lambda i: f"=IFERROR({P('REV', i)}/{avg('TR', i)},0)", '0.00'),
    ("Trade payables turnover ratio", "Net credit purchases", "Average trade payables",
     lambda i: f"=IFERROR({P('PUR', i)}/(({Bf('TP_MSME', CUR[i])}+{Bf('TP_OTH', CUR[i])}+{Bf('TP_MSME', PRV[i])}+{Bf('TP_OTH', PRV[i])})/2),0)", '0.00'),
    ("Net capital turnover ratio", "Net sales", "Working capital", lambda i: f"=IFERROR({P('REV', i)}/({Bf('CA', CUR[i])}-{Bf('CL', CUR[i])}),0)", '0.00'),
    ("Net profit ratio", "Net profit after tax", "Net sales", lambda i: f"=IFERROR({P('PAT', i)}/{P('REV', i)},0)", PCT),
    ("Return on capital employed", "Earnings before interest and tax", "Capital employed (net worth + total debt + DTL)",
     lambda i: f"=IFERROR(({P('PBT', i)}+{P('FIN', i)})/({Bf('EQ', CUR[i])}+{debt(CUR[i])}+{Bf('DTL', CUR[i])}),0)", PCT),
    ("Return on investment", "Interest income on deposits", "Average fixed deposits",
     lambda i: f"=IFERROR({Nn('Interest income', i)}/(({Nn('Fixed deposits with banks', i)}+{Nn('Fixed deposits with banks', i, True)})/2),0)", PCT),
]
dscr_cells = []
for nm, num, den, fn, fm in ratios:
    w(wsY, f"C{r}", nm); w(wsY, f"D{r}", num); w(wsY, f"E{r}", den)
    for i in range(3):
        v = fn(i)
        if "{REPAY}" in v: dscr_cells.append((r, i))
        w(wsY, f"{'FGH'[i]}{r}", v, fmt=fm)
    w(wsY, f"I{r}", f"=IFERROR(F{r}/G{r}-1,0)", fmt=PCT)
    w(wsY, f"J{r}", f"=IFERROR(G{r}/H{r}-1,0)", fmt=PCT)
    w(wsY, f"K{r}", f'=IF(OR(ABS(I{r})>25%,ABS(J{r})>25%),"[Reason to be given by management]","")', fI)
    r += 1
r += 1
w(wsY, f"C{r}", "Principal repayment of long-term borrowings during the year (for DSCR)")
for col in "FGH": w(wsY, f"{col}{r}", 0, fBlue, ACC, fill=yel)
for rr, i in dscr_cells:
    c = wsY[f"{'FGH'[i]}{rr}"]; c.value = c.value.replace("{REPAY}", str(r))

# ================================================================== Capitalisation
wsK = sheet(CP, {"A": 1, "B": 5, "C": 60, "D": 20, "E": 20})
head(wsK, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", span="E")
r = 8
w(wsK, f"B{r}", f"={q(AN)}!B8+1", fB); w(wsK, f"C{r}", "Capitalisation statement", fB); r += 2
w(wsK, f"C{r}", "Particulars", fB, bd=bTB); w(wsK, f"D{r}", dtxt("d26", "Pre-issue as at\n"), fB, bd=bTB, al=wrapc)
w(wsK, f"E{r}", "Post-issue *", fB, bd=bTB, al=wrapc); wsK.row_dimensions[r].height = 30; r += 1
w(wsK, f"C{r}", "Borrowings", fB); r += 1
s1 = r; w(wsK, f"C{r}", "Short-term borrowings"); w(wsK, f"D{r}", "=" + Bf("STB", "D"), fmt=ACC); r += 1
s2 = r; w(wsK, f"C{r}", "Long-term borrowings (including current maturities)"); w(wsK, f"D{r}", "=" + Bf("LTB", "D"), fmt=ACC); r += 1
s3 = r; w(wsK, f"C{r}", "Total borrowings (A)", fB, bd=bT); w(wsK, f"D{r}", f"=D{s1}+D{s2}", fB, ACC, bd=bT); r += 2
w(wsK, f"C{r}", "Shareholders' funds", fB); r += 1
e1_ = r; w(wsK, f"C{r}", "Share capital"); w(wsK, f"D{r}", "=" + Bf("SC", "D"), fmt=ACC); r += 1
e2_ = r; w(wsK, f"C{r}", "Reserves and surplus"); w(wsK, f"D{r}", "=" + Bf("RS", "D"), fmt=ACC); r += 1
e3_ = r; w(wsK, f"C{r}", "Total shareholders' funds (B)", fB, bd=bT); w(wsK, f"D{r}", f"=D{e1_}+D{e2_}", fB, ACC, bd=bT); r += 2
w(wsK, f"C{r}", "Long-term borrowings / shareholders' funds"); w(wsK, f"D{r}", f"=IFERROR(D{s2}/D{e3_},0)", fmt='0.00'); r += 1
w(wsK, f"C{r}", "Total borrowings / shareholders' funds (A / B)"); w(wsK, f"D{r}", f"=IFERROR(D{s3}/D{e3_},0)", fmt='0.00'); r += 2
for rr in range(s1, r - 1):
    if wsK[f"D{rr}"].value is not None: w(wsK, f"E{rr}", "[ ]", al=Alignment(horizontal="center"))
w(wsK, f"C{r}", "* Post-issue figures will be updated on finalisation of the issue price and number of shares.", fI)

# ================================================================== Additional Disclosures
wsD = sheet(AD, {"A": 1, "B": 5, "C": 120})
head(wsD, "ANNEXURE VI - NOTES TO RESTATED SUMMARY STATEMENTS", span="C")
r = 8
w(wsD, f"B{r}", f"={q(CP)}!B8+1", fB); w(wsD, f"C{r}", "Statement of additional disclosures with respect to amendments to Schedule III", fB); r += 2
for t in [
    "(i) The title deeds of all immovable properties are held in the name of the Company.",
    "(ii) The Company has not revalued its Property, Plant and Equipment or intangible assets during the years.",
    "(iii) The Company has not granted any loans or advances in the nature of loans to promoters, directors, KMPs or related parties.",
    "(iv) No proceedings have been initiated or are pending against the Company for holding any benami property under the Benami Transactions (Prohibition) Act, 1988.",
    "(v) The Company has borrowings from banks on the basis of security of current assets. Quarterly returns / statements filed with the banks are in agreement with the books of account, except as disclosed separately.",
    "(vi) The Company has not been declared a wilful defaulter by any bank or financial institution or other lender.",
    "(vii) The Company does not have any transactions with companies struck off under Section 248 of the Companies Act, 2013.",
    "(viii) There are no charges or satisfaction yet to be registered with the Registrar of Companies beyond the statutory period.",
    "(ix) The Company does not have any subsidiaries and hence compliance with the number of layers prescribed under Section 2(87) is not applicable.",
    "(x) Analytical ratios are disclosed separately in this Annexure.",
    "(xi) The Company has not advanced or loaned or invested funds to any other person or entity with the understanding that the intermediary shall lend or invest in other persons identified by the Company (Ultimate Beneficiaries).",
    "(xii) The Company has not surrendered or disclosed any income in the tax assessments under the Income-tax Act, 1961 that has not been recorded in the books of account.",
    "(xiii) The Company has not traded or invested in crypto currency or virtual currency during the years.",
    "(xiv) Corporate Social Responsibility under Section 135 is not applicable to the Company for the years presented."]:
    w(wsD, f"C{r}", t, al=just); wsD.row_dimensions[r].height = 28 if len(t) > 120 else 15; r += 1

# ================================================================== SOA_5
wsV = sheet(SA, {"A": 1, "B": 5, "C": 66, "D": 8, "E": 16.8, "F": 16.8, "G": 16.8})
head(wsV, "ANNEXURE V - STATEMENT OF ADJUSTMENTS TO AUDITED FINANCIAL STATEMENTS", span="G")
r = 8
w(wsV, f"B{r}", 1, fB); w(wsV, f"C{r}", "Reconciliation of restated profits", fB); r += 2
w(wsV, f"C{r}", "Particulars", fB, bd=bTB); w(wsV, f"D{r}", "Note", fB, bd=bTB, al=wrapc)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsV, f"{'EFG'[i]}{r}", dtxt(k, "For the year ended\n"), fB, bd=bTB, al=wrapc)
wsV.row_dimensions[r].height = 30
r += 1
pb = r
w(wsV, f"C{r}", "Net profit / (loss) after tax as per audited financial statements")
for i, col in enumerate("EFG"): w(wsV, f"{col}{r}", "=" + nref("__bookpat", "DEF"[i]), fmt=ACC)
r += 1
w(wsV, f"C{r}", "Adjusted for:", fI); r += 1


def pl_adj(i, label=None):
    fy = FYL[i]
    crit = (f'{AR("C")},"{label}",' if label else f'{AR("D")},"P&L",') + f'{AR("A")},">"&{md[START[fy]]},{AR("A")},"<="&{md[END[fy]]}'
    return f"(SUMIFS({AR('E')},{crit})-SUMIFS({AR('F')},{crit}))"


fa_ = r
w(wsV, f"C{r}", "Gratuity expense"); w(wsV, f"D{r}", "A", al=Alignment(horizontal="center"))
for i, col in enumerate("EFG"): w(wsV, f"{col}{r}", "=-" + pl_adj(i, "Gratuity"), fmt=ACC)
r += 1
w(wsV, f"C{r}", "Deferred tax"); w(wsV, f"D{r}", "B", al=Alignment(horizontal="center"))
for i, col in enumerate("EFG"): w(wsV, f"{col}{r}", "=-" + pl_adj(i, "Deferred tax"), fmt=ACC)
r += 1
w(wsV, f"C{r}", "Other adjustments")
for i, col in enumerate("EFG"): w(wsV, f"{col}{r}", f"=-{pl_adj(i)}-{col}{r-2}-{col}{r-1}", fmt=ACC)
r += 1
na_ = r
w(wsV, f"C{r}", "Net adjustment in Statement of Profit and Loss", fB, bd=bT)
for col in "EFG": w(wsV, f"{col}{r}", f"=SUM({col}{fa_}:{col}{r-1})", fB, ACC, bd=bT)
r += 1
w(wsV, f"C{r}", "Net profit / (loss) after tax as per Restated Summary Statements", fB, bd=bTot)
for i, col in enumerate("EFG"): w(wsV, f"{col}{r}", f"={col}{pb}+{col}{na_}", fB, ACC, bd=bTot)
for i, col in enumerate("EFG"): w(wsV, f"{'HIJ'[i]}{r}", f"=ROUND({col}{r}-{P('PAT', i)},2)", fS, ACC)
r += 2
w(wsV, f"C{r}", "Notes for reconciliation of profit:", fB); r += 1
for t in ["A) Gratuity: The Company recognised its gratuity liability for the first time in FY 2024-25 and debited the portion relating to earlier years directly to Reserves and "
          "surplus as a prior period item. In the Restated Summary Statements the liability has been recognised in the years to which it relates, with the portion relating to "
          "periods up to March 31, 2023 adjusted against opening reserves as at April 1, 2023.",
          "B) Deferred tax: Deferred tax has been recognised on the above timing difference in the respective years and reversed in FY 2024-25, when the audited financial "
          "statements carry the full liability."]:
    w(wsV, f"C{r}", t, al=just); wsV.merge_cells(f"C{r}:G{r}"); wsV.row_dimensions[r].height = 44; r += 1
r += 1
w(wsV, f"B{r}", 2, fB); w(wsV, f"C{r}", "Reconciliation of restated shareholders' funds", fB); r += 2
w(wsV, f"C{r}", "Particulars", fB, bd=bTB); w(wsV, f"D{r}", "Note", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsV, f"{'EFG'[i]}{r}", dtxt(k), fB, bd=bTB, al=wrapc)
wsV.row_dimensions[r].height = 30
r += 1
s_a = r
w(wsV, f"C{r}", "Shareholders' funds as per audited financial statements (a)")
for i, col in enumerate("EFG"):
    w(wsV, f"{col}{r}", f"={tref('SC', 'DEF'[i])}+{q(NE)}!{'DEF'[i]}{RSR['bookbal']}", fmt=ACC)
r += 1
w(wsV, f"C{r}", "Adjustments:", fI); r += 1
s_o = r
w(wsV, f"C{r}", "Opening balance of restatement adjustments")
for i, col in enumerate("EFG"):
    pc = "EFG"[i]
    w(wsV, f"{col}{r}", f"={q(NE)}!{pc}{RSR['close']}-{q(NE)}!{pc}{RSR['bookbal']}", fmt=ACC)
r += 1
w(wsV, f"C{r}", "Change in Statement of Profit and Loss (as above)")
for col in "EFG": w(wsV, f"{col}{r}", f"={col}{na_}", fmt=ACC)
r += 1
w(wsV, f"C{r}", "Adjustments routed directly through reserves"); w(wsV, f"D{r}", "A", al=Alignment(horizontal="center"))
for i, col in enumerate("EFG"):
    w(wsV, f"{col}{r}", f"={q(NE)}!{'DEF'[i]}{RSR['direct']}", fmt=ACC)
r += 1
s_b = r
w(wsV, f"C{r}", "Closing balance of restatement adjustments (b)", fB, bd=bT)
for col in "EFG": w(wsV, f"{col}{r}", f"=SUM({col}{s_o}:{col}{r-1})", fB, ACC, bd=bT)
r += 1
w(wsV, f"C{r}", "Restated shareholders' funds (a + b)", fB, bd=bTot)
for col in "EFG": w(wsV, f"{col}{r}", f"={col}{s_a}+{col}{s_b}", fB, ACC, bd=bTot)
for i, col in enumerate("EFG"): w(wsV, f"{'HIJ'[i]}{r}", f"=ROUND({col}{r}-{Bf('EQ', 'DEF'[i])},2)", fS, ACC)
r += 2
w(wsV, f"B{r}", 3, fB); w(wsV, f"C{r}", "Material regrouping", fB); r += 1
w(wsV, f"C{r}", "Appropriate regroupings have been made in the Restated Summary Statements, wherever required, by reclassification of the corresponding items of income, "
                "expenses, assets, liabilities and cash flows, in order to bring them in line with the accounting policies and classification adopted in the latest audited "
                "financial statements and the requirements of the SEBI ICDR Regulations. These include presentation of credit balances of trade receivables as advance from "
                "customers and debit balances of trade payables as advance to suppliers, group health insurance under employee benefits expense, credit card dues under other "
                "current liabilities and the gratuity provision under long-term provisions.", al=just)
wsV.merge_cells(f"C{r}:G{r}"); wsV.row_dimensions[r].height = 72; r += 2
w(wsV, f"B{r}", 4, fB); w(wsV, f"C{r}", "Non-adjusting items", fB); r += 1
w(wsV, f"C{r}", "There are no audit qualifications in the auditor's reports for the years ended March 31, 2026, March 31, 2025 and March 31, 2024 which require any adjustment "
                "in the Restated Summary Statements.", al=just)
wsV.merge_cells(f"C{r}:G{r}"); wsV.row_dimensions[r].height = 30

# ================================================================== source schedules: Shareholder Register, ageing
wsH = sheet(SR, {"A": 4, "B": 40, "C": 16, "D": 16, "E": 16})
w(wsH, "B2", "Register of members - number of equity shares held", fB)
w(wsH, "B4", "Name of shareholder", fB, bd=bTB)
for i, k in enumerate(["d26", "d25", "d24"]):
    w(wsH, f"{'CDE'[i]}4", f'=TEXT({md[k]},"dd-mm-yyyy")', fB, bd=bTB, al=wrapc)
holders = [("C K Lenin Christopher", 5880, 6000, 6000), ("W Vinolin Antony Joans", 5909, 9000, 9000), ("P Suresh Kumar", 2940, 0, 0),
           ("Akshaya Sathiyamoorthy", 14, 0, 0), ("Anoop Krishnan", 40, 0, 0), ("Ashwin", 14, 0, 0), ("Brittoraj L", 18, 0, 0), ("Dipu Ignatius", 10, 0, 0),
           ("Lavanya L", 14, 0, 0), ("Mervinraj Jeyaraj", 50, 0, 0), ("Muthu Lakshmi", 19, 0, 0), ("Palavesam Sivanupandian", 14, 0, 0),
           ("Rajendran Rajasekar", 14, 0, 0), ("Reyaz Ahamad", 10, 0, 0), ("R Saravanan", 14, 0, 0), ("Sankaranarayanan P", 30, 0, 0), ("Sunny Ramakrishna Pillai", 10, 0, 0)]
r = 5
for h in holders:
    w(wsH, f"B{r}", h[0])
    for i in range(3): w(wsH, f"{'CDE'[i]}{r}", h[1 + i], fBlue, ACC0)
    r += 1
w(wsH, f"B{r}", "Total", fB, bd=bTot)
for col in "CDE": w(wsH, f"{col}{r}", f"=SUM({col}5:{col}{r-1})", fB, ACC0, bd=bTot)
w(wsH, f"B{r+2}", "Source: Register of members / audited financial statements (Note on share capital).", fI)

for nm, buckets, data in [
    (DRA, ["Less than 6 months", "6 months - 1 year", "1-2 years", "2-3 years", "More than 3 years"],
     {"d25": [[65266644.79, 489208.21, 15890811.0, 7882898.58, 0]], "d24": [[70060872.03, 3304735.23, 11443850.94, 269831.44, 0]]}),
    (CRA, ["Less than 1 year", "1-2 years", "2-3 years", "More than 3 years"],
     {"d25": [[24757260, 0, 0, 0], [75016030, 2966003.37, 0, 0]], "d24": [[28602288.28, 0, 0, 0], [31960423.51, 1417265.97, 0, 0]]})]:
    wsZ = sheet(nm, {"A": 4, "B": 40, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16})
    w(wsZ, "B2", f"{nm} schedule (amount in Rs.)", fB)
    rows_lbl = (["(i) Undisputed - considered good", "(ii) Undisputed - considered doubtful", "(iii) Disputed - considered good", "(iv) Disputed - considered doubtful"]
                if nm == DRA else ["(i) MSME", "(ii) Others", "(iii) Disputed dues - MSME", "(iv) Disputed dues - Others"])
    for k, dk in enumerate(["d26", "d25", "d24"]):
        base = 5 + 8 * k
        w(wsZ, f"B{base-2}", dtxt(dk, "As at "), fB)
        w(wsZ, f"B{base-1}", "Particulars", fB, bd=bTB)
        for i, h in enumerate(buckets): w(wsZ, f"{CL(3+i)}{base-1}", h, fB, bd=bTB, al=wrapc)
        for j, lbl in enumerate(rows_lbl):
            w(wsZ, f"B{base+j}", lbl)
            for i in range(len(buckets)):
                v = data.get(dk, [[0] * 5] * 4)
                val = v[j][i] if j < len(v) else 0
                w(wsZ, f"{CL(3+i)}{base+j}", val, fBlue, ACC, fill=yel if dk == "d26" else None)
    w(wsZ, "B28", "March 31, 2025 and March 31, 2024: as per the ageing schedule in the audited financial statements for FY 2024-25. "
                  "March 31, 2026: to be filled from the ageing report.", fI)

# ------------------------------------------------------------------ sheet order (reference style)
order = [MD, BS, PL, CF, SAP, SA, NE, NL, NP, NA, NPL, RPT, GC, TX, AR_, AN, CP, AD, "TB 2025-26", "TB 2024-25", "TB 2023-24", DRA, CRA, SR, ADJ]
wb._sheets = [wb[s] for s in order]
for s in [BS, PL, CF, SAP, SA, NE, NL, NP, NA, NPL, RPT, GC, TX, AR_, AN, CP, AD]:
    wsx = wb[s]
    wsx.page_setup.orientation = "portrait"; wsx.page_setup.fitToWidth = 1; wsx.page_setup.fitToHeight = 0
    wsx.sheet_properties.pageSetUpPr.fitToPage = True
wb[BS].print_area = f"B2:F{wsB.max_row}"
wb[PL].print_area = f"B2:F{wsP.max_row}"
wb.active = 1
wb.save(OUT)
print("saved")
