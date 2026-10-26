import openpyxl, sys, json
def sign(c):
    v=c.value
    if v in (None,''): return 0.0
    return -float(v) if 'Cr' in c.number_format else float(v)
def load(path):
    ws=openpyxl.load_workbook(path).active
    out=[]; primary=None; groups=[]
    for r in range(8,ws.max_row+1):
        a=ws.cell(r,1)
        if a.value in (None,'') or str(a.value).strip()=='Grand Total': continue
        name=str(a.value).strip()
        kind='P' if a.font.b else ('L' if a.font.i else 'G')
        v=[sign(ws.cell(r,2)),float(ws.cell(r,3).value or 0),float(ws.cell(r,4).value or 0),sign(ws.cell(r,5))]
        if name=='Profit & Loss A/c': kind='P'
        if kind=='P': primary=name; lastgrp=''
        if kind=='G': lastgrp=name
        out.append(dict(r=r,name=name,kind=kind,primary=primary,group=lastgrp if kind!='P' else '',v=v))
    import collections
    cnt=collections.Counter((x['kind'] if x['kind']!='G' else 'L',x['name'],tuple(round(a,2) for a in x['v'])) for x in out)
    for x in out:
        x['dup']= x['kind']!='P' and any(abs(a)>0 for a in x['v']) and cnt[('L' if x['kind']=='G' else x['kind'],x['name'],tuple(round(a,2) for a in x['v']))]>1
    return out
if __name__=='__main__':
    rows=load(sys.argv[1])
    json.dump(rows,open(sys.argv[2],'w'))
    import collections
    P={x['name']:x['v'] for x in rows if x['kind']=='P'}
    S=collections.defaultdict(lambda:[0,0,0,0])
    for x in rows:
        if x['kind']=='L':
            for i in range(4): S[x['primary']][i]+=x['v'][i]
    # primaries with no ledgers (leaf primary)
    tot=[0]*4
    for p,v in P.items():
        s=S.get(p)
        if s is None: s=v; print('   primary w/o ledgers:',p,v)
        d=[round(a-b,2) for a,b in zip(v,s)]
        if any(abs(x)>0.01 for x in d): print('   MISMATCH',p,d)
        for i in range(4): tot[i]+=s[i]
    print(sys.argv[1][-14:],'ledgers',sum(1 for x in rows if x['kind']=='L'),'TB total (op,dr,cr,cl):',[round(t,2) for t in tot])
